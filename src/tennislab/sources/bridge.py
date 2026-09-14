"""Convert the RES2026 acquisition into the two input formats the chain joins.

Stage ``bridge``. Ported from ``references/WTA02_models/bridge_res2026.py``, the superset
of the ``TIER01_models`` copy: every TIER01 line survives under ``tour = "ATP"`` (the
diff is the ``tour`` switch, the WTA readers and manifest, the ``unresolved_event_policy``,
the event-end dating and the infobox reader); nothing TIER01-specific had to be merged
back. What changed in the port: modules are imported by name (``tennislab.panel.join``
for MULTI01's qualified workbook reader and round vocabulary, ``tennislab.panel.crosswalk_v2``
for identity), the two event-name tables are package data beside this module, every
workspace path goes through ``resolve_under_root``, the composed tarball's gzip header
carries a fixed mtime so its bytes are a function of its members alone, and the
synthesized-row dating has only the event-end route (below).

The chain starts from two sources and nothing else:

  * ``tennislab.panel.archive_panel`` reads Sackmann annual match files out of the pinned
    mirror tar (``atp/atp_matches_<year>.csv``), hash-gated;
  * the join reads one tennis-data annual workbook per season out of an acquisition
    manifest, and ``prepare_panel`` keeps only archive rows that a *market* row resolves
    to -- so a match with no tennis-data row gets no panel row, no feature row and no
    forecast.

RES2026 delivered neither format. It delivered tennis-data workbooks that stop at
**2026-08-03** and Wikipedia draw pages parsed to a per-draw CSV schema of its own. This
program converts both into the two formats above, under the rules
``experiments/CONFIRM2026.design.md`` declares in advance.

Overlap rule. Where the Sackmann mirror already has a match **the mirror row wins.** The
bridge regenerates the row anyway, compares it, and drops it as a duplicate on
``(event code, round, unordered winner/loser id pair)``; every drop is written to
``quarantine/dropped_duplicates.csv`` with both keys side by side.

Identity. Names resolve through ``crosswalk_v2``. Only ``high`` and ``medium`` confidence
resolutions are accepted (``crosswalk_v2.ACCEPTED_CONFIDENCE``); ``low`` ones are
quarantined to ``quarantine/unmatched_players.csv`` and counted; nothing is guessed.

Event identity, surface, level and best-of are all **carried forward from the same
event's most recent mirror edition**, never invented. The round labels and the draw size
come from the acquired draw's own structure: the numbered page rounds are laid on
Sackmann's ladder ending at the round of 16, and every round's match count is checked
against what that ladder round must hold. An event with no mirror edition, or a draw
whose counts fit no ladder, is quarantined whole.

Dates. Wikipedia draw pages carry no per-match date. A Wikipedia-only match gets the
tournament week start as ``tourney_date`` (the event anchor), and its synthesized market
row -- from which ``prepare_panel`` takes the panel row's ``match_date`` -- is dated at
the EVENT'S END (WTA02 declared rule 4, after the Astra review's finding 1: the men's
chain's start-plus-round-order dates admitted the Canada 2026 final four days before it
was played). The end is the completion date the retained draw page's infobox ``Date``
row gives when that page is retained, and otherwise the last week start in the calendar
cell plus 7 days. Every match of a draw therefore carries the same date, no result of
the event is admitted before the event could have finished, and the row's
``market_date_basis`` is ``inferred_event_end`` -- a bound, never a reported date. The
archive's WTA02 copy still carried the older start-plus-round-order route as a default
of ``synthesize_market_rows``; its driver never took it, and this module drops it, so
the start-plus-round-order basis string is never written as a date basis (the tests
pin that it does not occur in this file). The round order survives only as provenance
(``round_order``). A tennis-data row keeps its own
``Date`` (``tennis_data_reported_date``).

RESERVED WINDOW. This program opens files containing 2025/2026 match outcomes. It
refuses to run unless ``--allow-reserved-years`` is passed **and** the configuration
carries ``reserved_release_acknowledged``.

WTA mode (``bridge.tour = "WTA"``). The same program with the tour-specific halves
switched: the tennis-data workbook is read and written through the WTA01 ``wta_sources``
readers (the ones the WTA join uses), the market-side event identity comes from
``wta_market_event_names.json`` instead of MULTI01's crosswalk and CH01's pins, the
synthesized row carries the ``WTA`` event number and ``Tier`` columns instead of
``ATP``/``Series``, the merged market manifest is WTAODDS01's document extended, and
the round-cover index also keys on the event's ``Location``.
``unresolved_event_policy = "quarantine"`` makes an in-scope RES2026 event with no
carried mirror edition a quarantined draw instead of an aborted run; ATP runs keep
``refuse``.
"""

from __future__ import annotations

import argparse
import csv
import datetime as dt
import gzip
import html as html_module
import io
import json
import os
import re
import sys
import tarfile
import unicodedata
import zipfile
from collections import Counter, defaultdict
from collections.abc import Iterable, Mapping, Sequence
from pathlib import Path
from typing import Any

import openpyxl

from tennislab.chain.common import (
    ChainError,
    atomic_csv,
    atomic_json,
    code_receipt,
    read_config,
    read_csv_rows,
    relative_to_root,
    require_hash,
    resolve_output_under_root,
    resolve_under_root,
    sha256,
    sha256_bytes,
    year_plan,
)
from tennislab.chronology.dating import DateBasis
from tennislab.panel import crosswalk_v2

# Package data.  The tennis-data `Tournament`/`Location`/`ATP` the join itself saw for
# each archive event code in the carried seasons, derived from the join's own output
# (`work/CONFIRM2026_joint04/revision6/derive_market_event_names.py` in the archive), and
# its WTA analogue: per mirror event code, the tennis-data (Location, Tournament, WTA
# number, round_alignment) of the latest accepted edition in the frozen WTA crosswalk.
MARKET_EVENT_NAMES_FILE = Path(__file__).with_name("market_event_names.json")
WTA_MARKET_EVENT_NAMES_FILE = Path(__file__).with_name("wta_market_event_names.json")
# Workspace data the bridge reads (all resolved under the workspace root).
MARKET_ACQUISITION_DIR = "work/MULTI01_market_acquisition"
MARKET_ACQUISITION_MANIFEST = f"{MARKET_ACQUISITION_DIR}/acquisition_manifest.json"
MARKET_HEADERS_BY_YEAR = f"{MARKET_ACQUISITION_DIR}/headers_by_year.json"
DEFAULT_MARKET_EVENT_CROSSWALK = "work/MULTI01_event_crosswalk/qualified_event_crosswalk.csv"
WTA_MARKET_MANIFEST = "data/manifests/WTAODDS01.json"
RESERVED_YEARS = (2025, 2026)
UNRESOLVED_EVENT_POLICIES = ("refuse", "quarantine")

SACKMANN_FIELDS = (
    "tourney_id", "tourney_name", "surface", "draw_size", "tourney_level", "tourney_date",
    "match_num", "winner_id", "winner_seed", "winner_entry", "winner_name", "winner_hand",
    "winner_ht", "winner_ioc", "winner_age", "loser_id", "loser_seed", "loser_entry",
    "loser_name", "loser_hand", "loser_ht", "loser_ioc", "loser_age", "score", "best_of",
    "round", "minutes", "w_ace", "w_df", "w_svpt", "w_1stIn", "w_1stWon", "w_2ndWon",
    "w_SvGms", "w_bpSaved", "w_bpFaced", "l_ace", "l_df", "l_svpt", "l_1stIn",
    "l_1stWon", "l_2ndWon", "l_SvGms", "l_bpSaved", "l_bpFaced", "winner_rank",
    "winner_rank_points", "loser_rank", "loser_rank_points",
)  # fmt: skip
# Written by the bridge; blank on every generated row, which is what makes
# `count_block_status = missing_all` and so skips the SR02 state update.
BLANK_ON_GENERATED = (
    "winner_seed", "winner_entry", "winner_hand", "winner_ht", "winner_ioc", "winner_age",
    "loser_seed", "loser_entry", "loser_hand", "loser_ht", "loser_ioc", "loser_age",
    "minutes", "w_ace", "w_df", "w_svpt", "w_1stIn", "w_1stWon", "w_2ndWon", "w_SvGms",
    "w_bpSaved", "w_bpFaced", "l_ace", "l_df", "l_svpt", "l_1stIn", "l_1stWon",
    "l_2ndWon", "l_SvGms", "l_bpSaved", "l_bpFaced", "winner_rank", "winner_rank_points",
    "loser_rank", "loser_rank_points",
)  # fmt: skip

UNMATCHED_FIELDS = (
    "season", "tour", "tournament", "round", "side", "source_name", "reason",
    "match_method", "confidence", "candidate_ids", "source_url",
)  # fmt: skip
DUPLICATE_FIELDS = (
    "season", "dedupe_key", "generated_tourney_id", "mirror_tourney_id",
    "generated_round", "mirror_round", "generated_winner_id", "mirror_winner_id",
    "generated_loser_id", "mirror_loser_id", "generated_score", "mirror_score",
    "generated_tourney_date", "mirror_tourney_date", "generated_surface", "mirror_surface",
    "generated_draw_size", "mirror_draw_size", "generated_tourney_level",
    "mirror_tourney_level", "generated_best_of", "mirror_best_of",
    "generated_tourney_name", "mirror_tourney_name", "keys_identical",
)  # fmt: skip
UNMAPPED_EVENT_FIELDS = (
    "season", "tour", "tournament", "normalized_tournament", "matches", "reason", "detail",
)  # fmt: skip
MARKET_PROVENANCE_FIELDS = (
    "market_season", "market_source_row", "route", "tourney_id", "round",
    "winner_id", "loser_id", "price_source", "market_date_basis", "market_date", "round_order",
)  # fmt: skip

ROUND_ALIASES = {
    "f": "F", "final": "F", "the final": "F", "finals": "F",
    "sf": "SF", "semifinal": "SF", "semifinals": "SF", "semi finals": "SF",
    "qf": "QF", "quarterfinal": "QF", "quarterfinals": "QF", "quarter finals": "QF",
    "rr": "RR", "round robin": "RR",
    "br": "BR", "bronze medal match": "BR",
}  # fmt: skip
MARKET_ROUND_CANDIDATES = (
    "1st Round", "2nd Round", "3rd Round", "4th Round",
    "Quarterfinals", "Semifinals", "The Final", "Round Robin",
)  # fmt: skip
MONTH_NAMES = (
    "January", "February", "March", "April", "May", "June",
    "July", "August", "September", "October", "November", "December",
)  # fmt: skip
MONTHS = {name.lower(): number for number, name in enumerate(MONTH_NAMES, start=1)}

# The synthesized-row date is the event's end, a bound; a tennis-data row keeps the
# date its workbook reports.  `tennislab.chronology.dating.DateBasis` carries the
# archive's own strings, so nothing written changes.
SYNTHESIZED_DATE_BASIS = DateBasis.INFERRED_EVENT_END.value
TENNIS_DATA_DATE_BASIS = DateBasis.REPORTED_MATCH_DATE.value
SYNTHESIZED_ROUTE = "synthesized_from_wikipedia_draw_dated_inferred_event_end"


def normalized(value: str) -> str:
    """Casefold, fold `_` and `-` to spaces, collapse whitespace.

    RES2026's parsed draw rows carry the article-derived tournament title with
    underscores (`US_Open`, `French_Open`, `National_Bank_Open`); the mirror writes
    `Us Open`.  `.split()` collapses runs of whitespace; the underscore fold is what
    reconciles the two.
    """
    text = (value or "").casefold().replace("_", " ").replace("-", " ")
    return " ".join(text.split())


def _ascii_fold(value: str) -> str:
    """Drop combining marks and the apostrophes a sanitized file name loses.

    Kept out of `normalized()` deliberately: `normalized` is also used on round labels
    and retirement/walkover status strings, and only the *event alias* lookup needs to
    reconcile `Libema` with `Libéma` or `Queen_s` with `Queen's`.
    """
    decomposed = unicodedata.normalize("NFKD", value or "")
    stripped = "".join(ch for ch in decomposed if not unicodedata.combining(ch))
    return normalized(stripped.replace("'", " ").replace("’", " "))


# ------------------------------------------------- RES2026 event names and aliases
#
# The event names below are read from `references/RES2026.md`'s tournament table (the
# acquisition report), NOT from the 2026 data files: the reserved window must not be
# opened to fix a name.  RES2026 writes the Wikipedia *article* title, which is often
# the current sponsor name, while the Sackmann mirror writes its own short event name.
# Each alias maps a RES2026 spelling to the mirror name(s) to try, in order; nothing is
# guessed at run time, and an event with no resolution is refused loudly by
# `event_resolution_table()` under the `refuse` policy.
RES2026_EVENTS: tuple[tuple[str, str], ...] = (
    ("ATP", "French Open"),
    ("ATP", "BOSS Open"),
    ("ATP", "Libema Open"),
    ("ATP", "Halle Open"),
    ("ATP", "Queen's Club Championships"),
    ("ATP", "Mallorca Championships"),
    ("ATP", "Eastbourne Open"),
    ("ATP", "Wimbledon Championships"),
    ("ATP", "Swedish Open"),
    ("ATP", "Swiss Open Gstaad"),
    ("ATP", "Croatia Open Umag"),
    ("ATP", "Generali Open Kitzbuhel"),
    ("ATP", "Estoril Open"),
    ("ATP", "Mubadala Citi DC Open"),
    ("ATP", "Los Cabos Open"),
    ("ATP", "National Bank Open"),
    ("ATP", "Cincinnati Open"),
    ("ATP", "Winston-Salem Open"),
    ("ATP", "US Open"),
    ("WTA", "French Open"),
    ("WTA", "Queen's Club Championships"),
    ("WTA", "Libema Open"),
    ("WTA", "Berlin Tennis Open"),
    ("WTA", "Nottingham Open"),
    ("WTA", "Bad Homburg Open"),
    ("WTA", "Eastbourne Open"),
    ("WTA", "Wimbledon Championships"),
    ("WTA", "Iasi Open"),
    ("WTA", "Athens Open"),
    ("WTA", "Hamburg Open"),
    ("WTA", "Prague Open"),
    ("WTA", "Mubadala Citi DC Open"),
    ("WTA", "Memphis Classic"),
    ("WTA", "National Bank Open"),
    ("WTA", "Cincinnati Open"),
    ("WTA", "Monterrey Open"),
    ("WTA", "US Open"),
)

EVENT_NAME_ALIASES: dict[str, tuple[str, ...]] = {
    "french open": ("Roland Garros",),
    "boss open": ("Stuttgart",),
    "boss open stuttgart": ("Stuttgart",),
    "libema open": ("s Hertogenbosch", "'s-Hertogenbosch", "Hertogenbosch"),
    # RES2026's `tournament` column carries the *sanitized* article title, in which a
    # non-ASCII letter became `_` rather than being dropped: `Libéma` is written
    # `Lib_ma` and `Kitzbühel` `Kitzb_hel`, which `normalized()` folds to "lib ma" and
    # "kitzb hel".  `_ascii_fold` cannot recover a letter the acquisition already
    # replaced, so the sanitized spellings are registered from the probe's own file names.
    "lib ma open": ("s Hertogenbosch", "'s-Hertogenbosch", "Hertogenbosch"),
    "generali open kitzb hel": ("Kitzbuhel",),
    "ia i open": ("Iasi",),
    "halle open": ("Halle",),
    "queen s club championships": ("Queen's Club",),
    "queens club championships": ("Queen's Club",),
    "mallorca championships": ("Mallorca",),
    "eastbourne open": ("Eastbourne",),
    "wimbledon championships": ("Wimbledon",),
    "swedish open": ("Bastad",),
    "swiss open gstaad": ("Gstaad",),
    "croatia open umag": ("Umag",),
    "generali open kitzbuhel": ("Kitzbuhel",),
    "estoril open": ("Estoril",),
    "mubadala citi dc open": ("Washington",),
    "mubadala dc open": ("Washington",),
    "los cabos open": ("Los Cabos",),
    # The mirror has written this event as Canada Masters throughout 2005-2024; the
    # host-city spellings are registered because the event alternates Toronto/Montreal
    # and a later mirror edition may name the host instead.
    "national bank open": (
        "Canada Masters",
        "Toronto Masters",
        "Montreal Masters",
        "Toronto",
        "Montreal",
    ),  # fmt: skip
    # The two tours spell this event differently in the mirror -- ATP "Cincinnati
    # Masters" on every edition 2005-2024, WTA plain "Cincinnati" -- so both are
    # registered, ATP's first.  An ATP run resolves on the first candidate.
    "cincinnati open": ("Cincinnati Masters", "Cincinnati"),
    "winston salem open": ("Winston-Salem",),
    "us open": ("Us Open",),
    # WTA-only events, recorded from the same table.  An ATP run reports them as out of
    # scope rather than failing on them.
    "berlin tennis open": ("Berlin",),
    "nottingham open": ("Nottingham",),
    "bad homburg open": ("Bad Homburg",),
    "iasi open": ("Iasi",),
    "athens open": ("Athens",),
    "hamburg open": ("Hamburg",),
    "prague open": ("Prague",),
    "memphis classic": ("Memphis",),
    "monterrey open": ("Monterrey",),
}


def find_edition(
    index: Mapping[str, Mapping[str, Any]], tournament: str
) -> tuple[Mapping[str, Any] | None, str]:
    """Resolve a RES2026 tournament name to a carried mirror edition.

    Returns `(edition, how)`; `how` is `"normalized_name"` for a direct hit,
    `"alias:<mirror name>"` for a registered alias, `"ascii_folded_alias:<mirror name>"`
    when the RES2026 spelling only matches after dropping diacritics/apostrophes, and
    `""` when nothing resolves.
    """
    key = normalized(tournament)
    edition = index.get(f"name:{key}")
    if edition is not None:
        return edition, "normalized_name"
    for folded, label in ((key, "alias"), (_ascii_fold(tournament), "ascii_folded_alias")):
        for candidate in EVENT_NAME_ALIASES.get(folded, ()):
            edition = index.get(f"name:{normalized(candidate)}")
            if edition is not None:
                return edition, f"{label}:{candidate}"
    return None, ""


def event_resolution_table(
    index: Mapping[str, Mapping[str, Any]], tour: str, policy: str = "refuse"
) -> dict[str, Any]:
    """Every RES2026 event name and the mirror edition it resolves to.

    The 2024 round-trip renders *mirror* rows back into the Wikipedia schema, so it can
    never see a RES2026 spelling fail.  This check reads the names from
    `RES2026_EVENTS` -- the acquisition report's own list -- and resolves them against
    the carried mirror index, so it runs on exposed data and fails loudly (policy
    `refuse`) or lists the loss (policy `quarantine`) if any event of the configured
    tour has no edition.
    """
    rows: list[dict[str, Any]] = []
    for event_tour, name in RES2026_EVENTS:
        in_scope = event_tour.upper() == tour.upper()
        edition, how = find_edition(index, name)
        rows.append(
            {
                "tour": event_tour,
                "res2026_event_name": name,
                "res2026_event_name_underscored": name.replace(" ", "_"),
                "normalized": normalized(name),
                "in_scope_for_this_run": in_scope,
                "resolved": edition is not None,
                "resolved_via": how,
                "mirror_tourney_name": edition["tourney_name"] if edition else "",
                "mirror_tourney_id": edition["tourney_id"] if edition else "",
                "mirror_edition_season": edition["season"] if edition else "",
                "mirror_draw_size": edition["draw_size"] if edition else "",
                "mirror_tourney_level": edition["tourney_level"] if edition else "",
                "mirror_best_of": edition["best_of"] if edition else "",
            }
        )
    unresolved = [row for row in rows if row["in_scope_for_this_run"] and not row["resolved"]]
    report = {
        "tour": tour.upper(),
        "events_listed": len(rows),
        "events_in_scope": sum(1 for row in rows if row["in_scope_for_this_run"]),
        "events_resolved_in_scope": sum(
            1 for row in rows if row["in_scope_for_this_run"] and row["resolved"]
        ),
        "unresolved_in_scope": [row["res2026_event_name"] for row in unresolved],
        "resolutions": rows,
        "source": "references/RES2026.md tournament table (report, not data)",
        "unresolved_event_policy": policy,
        "status": (
            "PASS"
            if not unresolved
            else "PASS_WITH_UNRESOLVED_QUARANTINED"
            if policy == "quarantine"
            else "FAIL"
        ),
    }
    if policy not in UNRESOLVED_EVENT_POLICIES:
        raise ChainError(f"unknown unresolved_event_policy {policy!r}")
    if unresolved and policy == "refuse":
        raise ChainError(
            f"{len(unresolved)} RES2026 {tour.upper()} event name(s) resolve to no mirror "
            f"edition: {', '.join(row['res2026_event_name'] for row in unresolved)}; add the "
            "spelling to EVENT_NAME_ALIASES deliberately rather than widening normalized()"
        )
    return report


def numeric_round(label: str) -> int | None:
    text = normalized(label)
    match = re.fullmatch(r"r(?:ound)?\s*(\d+)", text)
    return int(match.group(1)) if match else None


def canonical_round(label: str) -> tuple[str, int | None]:
    """Return (canonical tail label or '', numeric round index or None)."""
    text = normalized(label)
    if text in ROUND_ALIASES:
        return ROUND_ALIASES[text], None
    index = numeric_round(text)
    if index is not None:
        return "", index
    raise ChainError(f"unknown round label {label!r}; add it to ROUND_ALIASES deliberately")


def month_number(token: str) -> int | None:
    """A month name or three-letter abbreviation -> 1-12, else None.

    A one- or two-letter token is not a month, and a bare number never reaches here
    because the tokenizer separates digits from letters.
    """
    name = (token or "").strip().rstrip(".").lower()
    if len(name) < 3 or not name.isalpha():
        return None
    return next((number for key, number in MONTHS.items() if key.startswith(name[:3])), None)


def week_starts(label: str) -> tuple[list[tuple[int, int]], int | None]:
    """Every week start in a calendar cell, in page order, plus a trailing year.

    RES2026 copies the Wikipedia calendar cell verbatim, and the two tours write it in
    opposite orders -- ATP `"8 Jun"`, WTA `"Jun 8"` -- while a **two-week** event
    concatenates both of its week starts with no separator at all: `"31 Aug7 Sep"`,
    `"Aug 31Sep 7"`.  The cell is tokenized into runs of digits and runs of letters,
    which is exactly the boundary the concatenation hides, and the tokens are then read
    as day/month pairs in either order.  A trailing four-digit year is accepted and
    returned separately.  Every token must belong to a pair or be that year: a cell
    carrying anything else is refused with the literal value rather than half-read.
    """
    text = " ".join((label or "").replace(",", " ").split())
    tokens = re.findall(r"\d+|[A-Za-z]+\.?", text)
    starts: list[tuple[int, int]] = []
    year: int | None = None
    position = 0
    while position < len(tokens):
        head = tokens[position]
        tail = tokens[position + 1] if position + 1 < len(tokens) else ""
        head_month, tail_month = month_number(head), month_number(tail)
        if head_month is not None and re.fullmatch(r"\d{1,2}", tail):
            starts.append((head_month, int(tail)))  # "Jun 8", WTA order
        elif tail_month is not None and re.fullmatch(r"\d{1,2}", head):
            starts.append((tail_month, int(head)))  # "8 Jun", ATP order
        elif starts and year is None and re.fullmatch(r"\d{4}", head) and not tail:
            year = int(head)
            position += 1
            continue
        else:
            raise ChainError(
                f"cannot parse week label {label!r}: {head!r} begins no day/month pair"
            )
        position += 2
    if not starts:
        raise ChainError(f"cannot parse week label {label!r}")
    return starts, year


def parse_week_label(label: str, season: int) -> dt.date:
    """The tournament's start date: the FIRST week start in the cell, in `season`.

    Reads `"31 Aug"`, `"31 August"`, `"Aug 31"`, `"2026-08-31"` and a concatenated
    two-week cell in either tour's order (`"31 Aug7 Sep"`, `"Aug 31Sep 7"`), with an
    optional trailing year.  Raises `ChainError` on a label it cannot read, which
    `generate_from_wikipedia` turns into a `week_label_unparseable` quarantine of that
    one event rather than an aborted run.
    """
    text = " ".join((label or "").replace(",", " ").split())
    try:
        parsed = dt.date.fromisoformat(text)
    except ValueError:
        pass
    else:
        if parsed.year != season:
            raise ChainError(f"week label {label!r} is not in season {season}")
        return parsed
    starts, year = week_starts(label)
    if year is not None and year != season:
        raise ChainError(f"week label {label!r} is not in season {season}")
    month, day = starts[0]
    try:
        return dt.date(season, month, day)
    except ValueError as error:
        raise ChainError(f"week label {label!r} is no date in season {season}: {error}") from error


INFOBOX_END_BASIS = "retained_page_infobox_date"
FALLBACK_END_BASIS = "last_week_start_plus_7_days"


def fallback_event_end(label: str, season: int) -> dt.date:
    """The last week start in the calendar cell plus 7 days, in `season`.

    On the real 2024 WTA calendar the Toronto and Cincinnati finals were played on the
    Monday after their week (start + 7), so start + 6 would date two real finals a day
    early.
    """
    starts, year = week_starts(label)
    if year is not None and year != season:
        raise ChainError(f"week label {label!r} is not in season {season}")
    month, day = starts[-1]
    return dt.date(season, month, day) + dt.timedelta(days=7)


def infobox_completion_date(html_path: Path, season: int) -> tuple[dt.date | None, str, str]:
    """(date, text, sha256) from the retained draw page's infobox `Date` row, or None.

    Reads the first `<th>Date</th><td>...</td>` pair of the page's infobox and takes the
    LAST day/month/year in that cell (`September 12, 2026`, `29 August 2026`, or a
    range's second date).  Nothing else on the page is read.  A cell that yields no date
    in `season` returns None and the caller falls back to the calendar-cell rule.
    """
    payload = html_path.read_bytes()
    text = payload.decode("utf-8", errors="replace")
    start = text.find("infobox")
    if start < 0:
        return None, "", sha256_bytes(payload)
    segment = text[start : start + 20000]
    match = re.search(r"<th[^>]*>\s*Date\s*</th>\s*<td[^>]*>(.*?)</td>", segment, re.S)
    if match is None:
        return None, "", sha256_bytes(payload)
    cell = " ".join(html_module.unescape(re.sub(r"<[^>]+>", " ", match.group(1))).split())
    found: list[dt.date] = []
    for day, month, year in re.findall(r"(\d{1,2})\s+([A-Za-z]+)\.?\s+(\d{4})", cell):
        number = month_number(month)
        if number is not None:
            try:
                found.append(dt.date(int(year), number, int(day)))
            except ValueError:
                pass
    for month, day, year in re.findall(r"([A-Za-z]+)\.?\s+(\d{1,2}),\s*(\d{4})", cell):
        number = month_number(month)
        if number is not None:
            try:
                found.append(dt.date(int(year), number, int(day)))
            except ValueError:
                pass
    found = [item for item in found if item.year == season]
    if not found:
        return None, cell, sha256_bytes(payload)
    return max(found), cell, sha256_bytes(payload)


def normalize_score(score: str, status: str) -> str:
    """Render a Wikipedia score in Sackmann's convention.

    `archive_panel.classify_status` reads the score alone: `W/O` is a walkover, a
    standalone `RET` token is a retirement, and any other letter makes the status
    `unknown` -- which the panel builder then excludes.  Wikipedia writes `retired` in
    words and renders an abandoned set as `0-0`, so both are rewritten here.
    """
    text = " ".join((score or "").replace(",", " ").split())
    state = normalized(status)
    if state in {"walkover", "w/o", "wo"} or "w/o" in normalized(text):
        return "W/O"
    body = re.sub(
        r"\b(?:retired|ret\.?|abandoned|abd|abn|defaulted|default|def\.?)\b",
        "",
        text,
        flags=re.IGNORECASE,
    )
    # Dash folding must precede the 0-0 filter: Wikipedia writes en/em dashes, so an
    # abandoned set arrives as "0-0" with a non-ASCII hyphen and would survive an
    # ASCII-only comparison. The 2024 self-test caught exactly that ordering.
    body = body.replace("−", "-").replace("–", "-").replace("—", "-")
    body = " ".join(token for token in body.split() if token != "0-0")
    if state in {"retired", "ret"}:
        return f"{body} RET".strip()
    if state in {"default", "defaulted", "disqualified"}:
        return f"{body} DEF".strip()
    if re.search(r"[A-Za-z]", body):
        raise ChainError(f"score {score!r} with status {status!r} still carries letters: {body!r}")
    return body


# ------------------------------------------------------------------ mirror reading


def _guard(year: int, allow_reserved: bool) -> None:
    if year in RESERVED_YEARS and not allow_reserved:
        raise ChainError(
            f"season {year} is inside the reserved 2025/2026 outcome window; pass "
            "--allow-reserved-years and set reserved_release_acknowledged in the config"
        )


def read_mirror_annual(
    archive: Path, tar_root: str, tour: str, year: int, *, allow_reserved: bool
) -> tuple[bytes, tuple[str, ...], list[dict[str, str]]]:
    # outcome-history read: a season's winner/loser rows, composed into the archive
    # and used for dedupe; nothing is scored.
    _guard(year, allow_reserved)
    member = f"{tar_root}/{tour.lower()}/{tour.lower()}_matches_{year}.csv"
    with tarfile.open(archive, "r:gz") as handle:
        extracted = handle.extractfile(member)
        if extracted is None:
            raise ChainError(f"mirror member missing: {member}")
        payload = extracted.read()
    reader = csv.DictReader(io.StringIO(payload.decode("utf-8-sig"), newline=""))
    header = tuple(reader.fieldnames or ())
    if header != SACKMANN_FIELDS:
        raise ChainError(f"mirror annual schema differs in {member}")
    return payload, header, [dict(row) for row in reader]


def event_index(
    archive: Path, tar_root: str, tour: str, years: Sequence[int], *, allow_reserved: bool
) -> dict[str, dict[str, Any]]:
    """Latest mirror edition per event, keyed by normalized tournament name and by code.

    Carries everything a generated row cannot invent: the Sackmann event code, name,
    surface, draw size, level, best-of and the ordered round labels with their match
    counts.  Later seasons overwrite earlier ones, so the carried edition is the most
    recent the mirror holds.
    """
    index: dict[str, dict[str, Any]] = {}
    for year in sorted(years):
        _guard(year, allow_reserved)
        _payload, _header, rows = read_mirror_annual(
            archive, tar_root, tour, year, allow_reserved=allow_reserved
        )
        editions: dict[str, dict[str, Any]] = {}
        for row in rows:
            tourney_id = row["tourney_id"].strip()
            if "-" not in tourney_id:
                continue
            code = tourney_id.rsplit("-", 1)[1]
            entry = editions.setdefault(
                code,
                {
                    "code": code,
                    "season": year,
                    "tourney_id": tourney_id,
                    "tourney_name": row["tourney_name"].strip(),
                    "surface": row["surface"].strip(),
                    "draw_size": row["draw_size"].strip(),
                    "tourney_level": row["tourney_level"].strip(),
                    "best_of": Counter(),
                    "rounds": Counter(),
                    "round_first_seen": {},
                    "max_match_num": 0,
                },
            )
            entry["best_of"][row["best_of"].strip()] += 1
            label = row["round"].strip()
            entry["rounds"][label] += 1
            entry["round_first_seen"].setdefault(label, len(entry["round_first_seen"]))
            try:
                entry["max_match_num"] = max(entry["max_match_num"], int(row["match_num"]))
            except ValueError:
                pass
        for code, entry in editions.items():
            entry["best_of"] = entry["best_of"].most_common(1)[0][0]
            entry["ordered_rounds"] = _draw_order(entry["rounds"])
            entry["rounds"] = dict(entry["rounds"])
            entry.pop("round_first_seen")
            index[code] = entry
            index[f"name:{normalized(entry['tourney_name'])}"] = entry
    return index


def _draw_order(rounds: Mapping[str, int]) -> list[str]:
    """Sackmann round labels in draw order: R<largest> .. R16, then QF, SF, F."""
    numbered = sorted(
        (int(label[1:]) for label in rounds if re.fullmatch(r"R\d+", label)), reverse=True
    )
    tail = [label for label in ("QF", "SF", "F") if label in rounds]
    others = sorted(
        label
        for label in rounds
        if not re.fullmatch(r"R\d+", label) and label not in ("QF", "SF", "F")
    )
    return [f"R{value}" for value in numbered] + tail + others


def align_rounds(
    draw_rounds: Mapping[str, int], edition: Mapping[str, Any]
) -> tuple[dict[str, str], str]:
    """Map the draw's round labels onto the mirror edition's, or explain the refusal.

    Numbered Wikipedia rounds (`R1`, `R2`, ...) are aligned positionally against the
    edition's leading `R<n>` labels; `QF`/`SF`/`F` map to themselves.  The alignment is
    accepted only when every aligned pair has the same match count, so a draw that grew
    or shrank relative to the carried edition is refused instead of silently relabelled.
    """
    numbered = sorted(
        ((numeric_round(label), label) for label in draw_rounds if numeric_round(label)),
        key=lambda item: item[0],
    )
    tails = [label for label in draw_rounds if numeric_round(label) is None]
    canonical_tails = {}
    for label in tails:
        head, index = canonical_round(label)
        if index is not None:
            raise ChainError(f"round {label!r} classified both ways")
        canonical_tails[label] = head
    edition_numbered = [
        label for label in edition["ordered_rounds"] if re.fullmatch(r"R\d+", label)
    ]
    if len(numbered) != len(edition_numbered):
        return {}, (f"numbered_round_count_differs:{len(numbered)}_vs_{len(edition_numbered)}")
    mapping: dict[str, str] = {}
    for (_index, label), target in zip(numbered, edition_numbered, strict=True):
        if draw_rounds[label] != edition["rounds"].get(target):
            return {}, (
                f"round_match_count_differs:{label}={draw_rounds[label]}"
                f"_vs_{target}={edition['rounds'].get(target)}"
            )
        mapping[label] = target
    for label, head in canonical_tails.items():
        if head not in edition["rounds"]:
            return {}, f"edition_lacks_round:{head}"
        if draw_rounds[label] != edition["rounds"][head]:
            return {}, (
                f"round_match_count_differs:{label}={draw_rounds[label]}"
                f"_vs_{head}={edition['rounds'][head]}"
            )
        mapping[label] = head
    return mapping, ""


# A Sackmann round label is the number of players *remaining* in that round, so it is
# fixed by counting rounds back from the final -- not by the page's own numbering, and
# not by the carried edition's match counts.  `LADDER` is that vocabulary in order from
# the final outwards and `REMAINING` the players remaining, so a round's own nominal
# match count is half of it.  Checked against the pinned mirror: over all 1,252
# ladder-shaped ATP editions 2005-2024 this derivation reproduces the mirror's own round
# label on every round, and its bracket top equals the mirror's `draw_size` bracket.
LADDER = ("F", "SF", "QF", "R16", "R32", "R64", "R128", "R256")
REMAINING = {"F": 2, "SF": 4, "QF": 8, "R16": 16, "R32": 32, "R64": 64, "R128": 128, "R256": 256}


def bracket_top(draw_size: Any) -> int | None:
    """The power-of-two bracket a recorded `draw_size` sits in, or None."""
    text = str(draw_size or "").strip()
    if not text.isdigit() or int(text) < 2:
        return None
    return 1 << (int(text) - 1).bit_length()


def derive_draw_structure(
    draw_rounds: Mapping[str, int], edition: Mapping[str, Any]
) -> tuple[dict[str, Any] | None, str]:
    """Label the acquired draw's rounds from the draw's own structure.

    The draw's own round column and order decide the labels: the numbered page rounds
    are laid on the ladder ending at the round of 16 (the round before the
    quarterfinals), the page's `QF`/`SF`/`F` label themselves, and every round's match
    count is then checked against what that ladder round must hold -- exactly, except
    the draw's first round, which is short by its byes.  A draw whose counts do not fit
    the ladder is refused; a draw that merely differs from the carried edition is not.
    `draw_size` stays the carried edition's own value while the bracket is unchanged
    (the mirror's recorded value is not a consistent convention: 2024 records 64 for a
    56-draw Cincinnati and 128 for a 96-draw Indian Wells) and becomes the acquired
    draw's bracket when the bracket itself changed.  Missing later rounds are reported
    as pending, not invented.
    """
    numbered = sorted(
        ((numeric_round(label), label) for label in draw_rounds if numeric_round(label)),
        key=lambda item: item[0],
    )
    tails: dict[str, str] = {}
    for label in draw_rounds:
        if numeric_round(label) is not None:
            continue
        head, index = canonical_round(label)
        if index is not None:
            raise ChainError(f"round {label!r} classified both ways")
        tails[label] = head
    outside = sorted({head for head in tails.values() if head not in LADDER})
    if outside:
        return None, f"non_ladder_round_label:{','.join(outside)}"

    index_of: dict[str, int] = {}
    for offset, (_position, label) in enumerate(reversed(numbered)):
        index_of[label] = LADDER.index("R16") + offset
    for label, head in tails.items():
        index_of[label] = LADDER.index(head)
    if not index_of:
        return None, "draw_has_no_rounds"
    if max(index_of.values()) >= len(LADDER):
        return None, f"draw_deeper_than_the_round_vocabulary:{len(numbered)}_numbered_rounds"
    present = sorted(index_of.values())
    if len(set(present)) != len(present):
        return None, "two_page_rounds_map_to_one_ladder_round"
    if present != list(range(present[0], present[-1] + 1)):
        return None, f"rounds_not_contiguous:{','.join(LADDER[index] for index in present)}"

    first = present[-1]
    for label, index in index_of.items():
        nominal = REMAINING[LADDER[index]] // 2
        count = draw_rounds[label]
        if index == first:
            if not 0 < count <= nominal:
                return None, (
                    f"first_round_matches_{count}_outside_1_{nominal}_for_{LADDER[index]}"
                )
        elif count != nominal:
            return None, (f"round_match_count_differs:{LADDER[index]}={count}_expected_{nominal}")

    mapping = {label: LADDER[index] for label, index in index_of.items()}
    page_order = [
        label
        for _index, label in sorted(
            ((index, label) for label, index in index_of.items()), reverse=True
        )
    ]
    top = REMAINING[LADDER[first]]
    mirror_size = str(edition.get("draw_size", "")).strip()
    warnings: list[str] = []
    if bracket_top(mirror_size) == top:
        draw_size, draw_size_basis = mirror_size, "carried_edition"
    else:
        draw_size, draw_size_basis = str(top), "acquired_draw"
        warnings.append(
            f"draw_size_from_acquired_draw:{draw_size}_carried_edition:{mirror_size or 'unknown'}"
        )
    pending = [LADDER[index] for index in range(present[0] - 1, -1, -1)]
    if pending:
        warnings.append(f"rounds_pending:{','.join(pending)}")
    edition_matches = sum(int(value) for value in edition.get("rounds", {}).values())
    matches = sum(int(value) for value in draw_rounds.values())
    if sorted(mapping.values()) != sorted(edition.get("rounds", {})) or matches != edition_matches:
        warnings.append(
            f"round_structure_differs_from_carried_edition:{matches}_matches_vs_{edition_matches}"
        )
    return {
        "basis": "acquired_draw_structure",
        "mapping": mapping,
        "page_round_order": page_order,
        "rounds_generated": [mapping[label] for label in page_order],
        "rounds_pending": pending,
        "draw_size": draw_size,
        "draw_size_basis": draw_size_basis,
        "carried_edition_draw_size": mirror_size,
        "matches": matches,
        "carried_edition_matches": edition_matches,
        "warnings": warnings,
    }, ""


def draw_structure(
    draw_rounds: Mapping[str, int], edition: Mapping[str, Any]
) -> tuple[dict[str, Any] | None, str]:
    """The acquired draw's own structure, or the carried edition's alignment.

    A round-robin or bronze-match draw has no ladder to count back along, so for those
    `align_rounds`' count alignment against the carried edition remains the rule.
    """
    structure, refusal = derive_draw_structure(draw_rounds, edition)
    if structure is not None:
        return structure, ""
    if not refusal.startswith("non_ladder_round_label"):
        return None, refusal
    mapping, edition_refusal = align_rounds(draw_rounds, edition)
    if edition_refusal:
        return None, f"{refusal};{edition_refusal}"
    order = {label: position for position, label in enumerate(edition["ordered_rounds"])}
    page_order = sorted(mapping, key=lambda label: order.get(mapping[label], 99))
    return {
        "basis": "carried_edition_alignment",
        "mapping": mapping,
        "page_round_order": page_order,
        "rounds_generated": [mapping[label] for label in page_order],
        "rounds_pending": [],
        "draw_size": str(edition.get("draw_size", "")).strip(),
        "draw_size_basis": "carried_edition",
        "carried_edition_draw_size": str(edition.get("draw_size", "")).strip(),
        "matches": sum(int(value) for value in draw_rounds.values()),
        "carried_edition_matches": sum(int(value) for value in edition.get("rounds", {}).values()),
        "warnings": [],
    }, ""


# ------------------------------------------------------------------ source reading


def read_wikipedia(paths: Sequence[Path]) -> list[dict[str, str]]:
    # outcome-history read: draw-page winners and losers, converted to archive rows.
    rows: list[dict[str, str]] = []
    required = {"tour", "tournament", "round", "winner", "loser", "score", "status", "week_label"}
    for path in paths:
        header, parsed = read_csv_rows(path)
        missing = required - set(header)
        if missing:
            raise ChainError(f"{path} lacks Wikipedia draw columns {sorted(missing)}")
        for row in parsed:
            row["_source_path"] = str(path)
            rows.append(row)
    return rows


def join_module() -> Any:
    """MULTI01's qualified join, imported by name (the archive loaded it by path)."""
    from tennislab.panel import join

    return join


# The WTA workbook readers the WTA join uses, copied verbatim from the archive's
# `references/WTA01_event_map/wta_sources.py` (`_read_xlsx`, `_as_date`, `_blank`,
# `norm_name`).  That file belongs to the WTA join port (`tennislab.sources.wta_sources`);
# once it lands these four should become imports from it and this block go.


def _wta_strip_accents(value: str) -> str:
    return "".join(
        ch for ch in unicodedata.normalize("NFKD", value) if not unicodedata.combining(ch)
    )


def wta_norm_name(value: Any) -> str:
    """Loose fold of an event name: accents, punctuation, case, whitespace."""
    if value is None:
        return ""
    text = _wta_strip_accents(str(value)).lower()
    text = text.replace("&", " and ")
    text = re.sub(r"[^a-z0-9]+", " ", text)
    return " ".join(text.split())


def _wta_read_xlsx(path: Path) -> tuple[list[str], list[list[Any]]]:
    book = openpyxl.load_workbook(path, read_only=True, data_only=True)
    sheet = book[book.sheetnames[0]]
    it = sheet.iter_rows(values_only=True)
    header = [("" if v is None else str(v)).strip() for v in next(it)]
    rows = [list(r) for r in it]
    book.close()
    return header, rows


def _wta_as_date(value: Any) -> dt.date | None:
    if isinstance(value, dt.datetime):
        return value.date()
    if isinstance(value, dt.date):
        return value
    if isinstance(value, str) and value.strip():
        for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d/%m/%y", "%m/%d/%Y"):
            try:
                return dt.datetime.strptime(value.strip(), fmt).date()
            except ValueError:
                pass
    return None


def _wta_blank(row: Iterable[Any]) -> bool:
    return all(v is None or (isinstance(v, str) and v.strip() == "") for v in row)


class _WtaJoinShim:
    """The WTA chain does not use MULTI01's join; only its `normalized` fold is needed
    for the round-cover index, and the WTA sources' fold stands in for it."""

    @staticmethod
    def normalized(value: str) -> str:
        return wta_norm_name(value)


def read_wta_workbook(path: Path, season: int) -> tuple[list[str], list[dict[str, Any]]]:
    """Read a tennis-data WTA workbook exactly as the WTA join will.

    The `wta_sources` readers are the ones the WTA join uses, so the bridge sees the
    same header, the same cell values and the same dates.  Blank rows are skipped as
    the join skips them; `market_source_row` is the sheet row number.
    """
    # outcome-history read: workbook winner/loser columns, passed through unchanged.
    header, raw = _wta_read_xlsx(path)
    columns = [name for name in header if name]
    index = {name: position for position, name in enumerate(header) if name}
    rows: list[dict[str, Any]] = []
    for number, values in enumerate(raw, start=2):
        if _wta_blank(values):
            continue
        record: dict[str, Any] = {}
        for name, position in index.items():
            record[name] = values[position] if position < len(values) else None
        record.update(
            market_season=season,
            market_source_path=relative_to_root(path, label="tennis_data"),
            market_source_row=number,
            market_date=_wta_as_date(record.get("Date")),
        )
        rows.append(record)
    return columns, rows


def read_tennis_data(
    path: Path, season: int, tour: str = "ATP"
) -> tuple[list[str], list[dict[str, Any]]]:
    """Read a tennis-data annual workbook, or a RES2026 `parsed/*_window.csv` export.

    ATP workbooks go through the pinned `join.read_market_workbook`, so the
    canonicalization, date handling and formula refusal are the ones MULTI01 qualified;
    a WTA workbook goes through the `wta_sources` readers instead.  The RES2026 CSV
    export carries the same tennis-data columns plus four provenance columns, which are
    dropped here and recorded in the provenance sidecar instead.
    """
    # outcome-history read: workbook winner/loser columns, passed through unchanged.
    if path.suffix.casefold() in {".xlsx", ".xls"}:
        if tour.upper() == "WTA":
            return read_wta_workbook(path, season)
        headers, rows = join_module().read_market_workbook(path, season)
        return list(headers), rows
    header, parsed = read_csv_rows(path)
    provenance = ("source", "source_url", "retrieved_at_utc", "archive_capture")
    columns = [name for name in header if name not in provenance]
    rows: list[dict[str, Any]] = []
    for number, row in enumerate(parsed, 2):
        record = {name: row[name] for name in columns}
        raw_date = row.get("Date", "")
        record.update(
            market_season=season,
            market_source_path=relative_to_root(path, label="tennis_data"),
            market_source_row=number,
            market_date=_parse_market_date(raw_date),
        )
        for name in provenance:
            record[f"_{name}"] = row.get(name, "")
        rows.append(record)
    return columns, rows


def _parse_market_date(raw: str) -> dt.date | None:
    text = (raw or "").strip()
    if not text:
        return None
    for pattern in ("%Y-%m-%d", "%d/%m/%Y", "%m/%d/%Y", "%Y-%m-%d %H:%M:%S"):
        try:
            return dt.datetime.strptime(text[: len(pattern) + 4], pattern).date()
        except ValueError:
            continue
    try:
        return dt.date.fromisoformat(text[:10])
    except ValueError as error:
        raise ChainError(f"cannot parse market date {raw!r}") from error


# ------------------------------------------------------------------ row generation


def dedupe_key(tourney_id: str, round_label: str, winner_id: str, loser_id: str) -> tuple[str, ...]:
    code = tourney_id.rsplit("-", 1)[1] if "-" in tourney_id else tourney_id
    pair = tuple(sorted((str(winner_id), str(loser_id))))
    return (code, round_label, *pair)


class Resolver:
    """Name -> id with the bridge's confidence policy, memoized, with a quarantine log."""

    def __init__(self, crosswalks: Mapping[str, Any]):
        self.crosswalks = crosswalks
        self.cache: dict[tuple[str, str], dict[str, Any]] = {}
        self.rejected: list[dict[str, Any]] = []

    def resolve(self, tour: str, name: str) -> dict[str, Any]:
        key = (tour.upper(), (name or "").strip())
        if key not in self.cache:
            crosswalk = self.crosswalks.get(tour.upper())
            if crosswalk is None:
                raise ChainError(f"no crosswalk built for tour {tour!r}")
            record = crosswalk.lookup(key[1])
            if record["matched"] and record["confidence"] not in crosswalk_v2.ACCEPTED_CONFIDENCE:
                record = {
                    **record,
                    "matched": False,
                    "reason": f"below_accepted_confidence_{record['match_method']}",
                    "candidate_ids": str(record.get("player_id") or ""),
                }
            self.cache[key] = record
        return self.cache[key]


def _unmapped(season: int, tour: str, tournament: str, draw: int, reason: str, detail: str):
    return {
        "season": season,
        "tour": tour,
        "tournament": tournament,
        "normalized_tournament": normalized(tournament),
        "matches": draw,
        "reason": reason,
        "detail": detail,
    }


def event_end(
    ordered_first: Mapping[str, str], season: int, anchor: dt.date, infobox_dir: Path | None
) -> tuple[dt.date, str, str, str]:
    """The draw's completion date and its basis: the retained page's infobox `Date`
    when `infobox_dir` holds the page (`<csv stem>.html`) and it is on or after the
    anchor, else the last week start plus 7 days."""
    end_date = fallback_event_end(ordered_first["week_label"], season)
    end_basis, infobox_text, infobox_sha = FALLBACK_END_BASIS, "", ""
    source_path = ordered_first.get("_source_path", "")
    if infobox_dir is not None and source_path:
        html_path = Path(infobox_dir) / (Path(source_path).stem + ".html")
        if html_path.is_file():
            parsed, infobox_text, infobox_sha = infobox_completion_date(html_path, season)
            if parsed is not None and parsed >= anchor:
                end_date, end_basis = parsed, INFOBOX_END_BASIS
            else:
                end_basis = FALLBACK_END_BASIS + ":infobox_unusable"
    return end_date, end_basis, infobox_text, infobox_sha


def generate_from_wikipedia(
    rows: Sequence[Mapping[str, str]],
    season: int,
    index: Mapping[str, Mapping[str, Any]],
    resolver: Resolver,
    infobox_dir: Path | None = None,
) -> tuple[
    list[dict[str, str]], list[dict[str, Any]], list[dict[str, Any]], dict[str, dict[str, Any]]
]:
    """Sackmann-format rows, the two quarantine tables, and the per-draw event resolution.

    Each draw's resolution record also carries `tournament_end_date` and its basis --
    the retained page's infobox completion date when `infobox_dir` holds the page, else
    the last week start plus 7 days -- which `synthesize_market_rows` dates every
    synthesized row of the draw at.
    """
    unmapped: list[dict[str, Any]] = []
    resolutions: dict[tuple[str, str], dict[str, Any]] = {}
    by_draw: dict[tuple[str, str], list[Mapping[str, str]]] = defaultdict(list)
    for row in rows:
        by_draw[(row["tour"].strip().upper(), row["tournament"].strip())].append(row)

    generated: list[dict[str, str]] = []
    for (tour, tournament), draw in sorted(by_draw.items()):
        edition, resolved_via = find_edition(index, tournament)
        if edition is None:
            unmapped.append(
                _unmapped(
                    season,
                    tour,
                    tournament,
                    len(draw),
                    "no_mirror_edition_for_event",
                    "no earlier mirror edition carries this tournament name, "
                    "directly or through a registered RES2026 alias",
                )
            )
            continue
        counts = Counter(row["round"].strip() for row in draw)
        try:
            structure, refusal = draw_structure(counts, edition)
        except ChainError as error:
            unmapped.append(
                _unmapped(season, tour, tournament, len(draw), "unknown_round_label", str(error))
            )
            continue
        if structure is None:
            unmapped.append(
                _unmapped(season, tour, tournament, len(draw), "round_structure_mismatch", refusal)
            )
            continue
        mapping = structure["mapping"]
        # The draw's own round order, so a round the carried edition does not have
        # (`R128` against a 56-draw edition) still sorts in its real position.
        # `match_num` is the position in that order, offset above the carried edition's
        # maximum so it cannot collide with a mirror row of the same event.
        order = {label: position for position, label in enumerate(structure["page_round_order"])}
        ordered = sorted(
            draw,
            key=lambda row: (
                order[row["round"].strip()],
                row.get("winner", ""),
                row.get("loser", ""),
            ),
        )
        # The week start is read here, before the draw is recorded as resolved, so a
        # calendar cell the parser cannot read quarantines this one event instead of
        # aborting the whole bridge.
        try:
            anchor = parse_week_label(ordered[0]["week_label"], season)
            end_date, end_basis, infobox_text, infobox_sha = event_end(
                ordered[0], season, anchor, infobox_dir
            )
        except ChainError as error:
            unmapped.append(
                _unmapped(season, tour, tournament, len(draw), "week_label_unparseable", str(error))
            )
            continue
        resolutions[(tour, tournament)] = {
            "resolved_via": resolved_via,
            "mirror_tourney_name": edition["tourney_name"],
            "mirror_tourney_id": edition["tourney_id"],
            "generated_tourney_id": f"{season}-{edition['code']}",
            "matches": len(draw),
            "tournament_end_date": end_date.isoformat(),
            "event_end_basis": end_basis,
            "infobox_date_text": infobox_text,
            "infobox_html_sha256": infobox_sha,
            "week_starts_in_cell": len(week_starts(ordered[0]["week_label"])[0]),
            # The draw's own structure, and every way it differs from the carried
            # edition, recorded per draw rather than quarantining the event.
            "round_labels_from": structure["basis"],
            "rounds_generated": structure["rounds_generated"],
            "rounds_pending": structure["rounds_pending"],
            "draw_size": structure["draw_size"],
            "draw_size_from": structure["draw_size_basis"],
            "carried_edition_draw_size": structure["carried_edition_draw_size"],
            "carried_edition_matches": structure["carried_edition_matches"],
            "warnings": structure["warnings"],
            "week_label": ordered[0]["week_label"],
            "tournament_start_date": anchor.isoformat(),
        }
        next_match_num = int(edition["max_match_num"]) + 1
        for row in ordered:
            winner_name = row.get("winner") or row.get("winner_display", "")
            loser_name = row.get("loser") or row.get("loser_display", "")
            winner = resolver.resolve(tour, winner_name)
            loser = resolver.resolve(tour, loser_name)
            rejects = [
                (side, name, record)
                for side, name, record in (
                    ("winner", winner_name, winner),
                    ("loser", loser_name, loser),
                )
                if not record["matched"]
            ]
            if rejects:
                for side, name, record in rejects:
                    resolver.rejected.append(
                        {
                            "season": season,
                            "tour": tour,
                            "tournament": tournament,
                            "round": row["round"],
                            "side": side,
                            "source_name": name,
                            "reason": str(record.get("reason", "")),
                            "match_method": str(record.get("match_method", "")),
                            "confidence": str(record.get("confidence", "")),
                            "candidate_ids": str(record.get("candidate_ids", "")),
                            "source_url": row.get("source_url", ""),
                        }
                    )
                continue
            record = dict.fromkeys(SACKMANN_FIELDS, "")
            record.update(
                {
                    "tourney_id": f"{season}-{edition['code']}",
                    "tourney_name": edition["tourney_name"],
                    "surface": edition["surface"],
                    "draw_size": structure["draw_size"],
                    "tourney_level": edition["tourney_level"],
                    "tourney_date": anchor.strftime("%Y%m%d"),
                    "match_num": str(next_match_num),
                    "winner_id": str(winner["player_id"]),
                    "winner_name": winner["sackmann_name"],
                    "loser_id": str(loser["player_id"]),
                    "loser_name": loser["sackmann_name"],
                    "score": normalize_score(row.get("score", ""), row.get("status", "")),
                    "best_of": edition["best_of"],
                    "round": mapping[row["round"].strip()],
                }
            )
            for field in BLANK_ON_GENERATED:
                record[field] = ""
            generated.append(record)
            next_match_num += 1
    return (
        generated,
        resolver.rejected,
        unmapped,
        {f"{tour}|{tournament}": record for (tour, tournament), record in resolutions.items()},
    )


def merge_annual(
    mirror_rows: Sequence[Mapping[str, str]],
    generated: Sequence[Mapping[str, str]],
    season: int,
) -> tuple[list[dict[str, str]], list[dict[str, Any]]]:
    """Mirror rows first and unchanged; generated rows only where the mirror has none."""
    mirror_by_key: dict[tuple[str, ...], Mapping[str, str]] = {}
    for row in mirror_rows:
        key = dedupe_key(row["tourney_id"], row["round"], row["winner_id"], row["loser_id"])
        mirror_by_key.setdefault(key, row)
    kept: list[dict[str, str]] = [dict(row) for row in mirror_rows]
    dropped: list[dict[str, Any]] = []
    seen_generated: set[tuple[str, ...]] = set()
    for row in generated:
        key = dedupe_key(row["tourney_id"], row["round"], row["winner_id"], row["loser_id"])
        mirror = mirror_by_key.get(key)
        if mirror is None and key in seen_generated:
            mirror = next(
                item
                for item in kept
                if dedupe_key(
                    item["tourney_id"], item["round"], item["winner_id"], item["loser_id"]
                )
                == key
            )
        if mirror is not None:
            compared = (
                "tourney_id", "round", "winner_id", "loser_id", "score", "tourney_date",
                "surface", "draw_size", "tourney_level", "best_of", "tourney_name",
            )  # fmt: skip
            dropped.append(
                {
                    "season": season,
                    "dedupe_key": "|".join(key),
                    "keys_identical": all(
                        str(row[field]) == str(mirror[field]) for field in compared
                    ),
                    **{f"generated_{field}": row[field] for field in compared},
                    **{f"mirror_{field}": mirror[field] for field in compared},
                }
            )
            continue
        seen_generated.add(key)
        kept.append(dict(row))
    return kept, dropped


# ------------------------------------------------------------------ market side


def round_order(round_label: str, draw_size: Any) -> int | None:
    """Rounds elapsed since the draw's first round: `R128` -> 0 ... `F` -> 6 in a 128.

    Provenance only.  The archive's first design dated a Wikipedia-only match at the
    tournament start plus this order; that put the Canada 2026 final four days before it
    was played, and the synthesized row is now dated at the event's end.  The order is
    still recorded beside the row (`round_order`) so the earlier approximation stays
    auditable, and it is never a date.
    """
    label = (round_label or "").strip()
    top = bracket_top(draw_size)
    remaining = REMAINING.get(label)
    if remaining is None and re.fullmatch(r"R\d+", label):
        remaining = int(label[1:])
    if top is None or remaining is None or remaining < 2 or remaining > top:
        return None
    return (top // remaining).bit_length() - 1


def market_event_identity(
    match: Mapping[str, str],
    market_events: Mapping[str, Mapping[str, str]],
    event_names: Mapping[str, Mapping[str, Any]],
    join: Any,
) -> tuple[dict[str, str], str, str]:
    """The `Tournament`/`Location`/`ATP` of one synthesized market row, and its basis.

    `join.event_evidence` accepts three strong paths.  `pinned_ch01_event` needs a CH01
    row and CH01 covers 2023-2024 only; `event_code_exact` needs the tennis-data `ATP`
    number to equal the Sackmann event code, which it never does.  The one path open to
    a synthesized row of a season nobody has pinned is `event_name_exact_normalized`, so
    the `Tournament` written here is the name that satisfies it: the carried edition's
    tennis-data spelling when that already normalizes to the archive name (`US Open`
    for `Us Open`), and otherwise the archive event's own `tourney_name`, which
    satisfies it by construction.  `Location` and the `ATP` number are carried from
    real tennis-data values wherever a carried season has them; neither participates in
    the join's event evidence.  Nothing is refused.
    """
    code = match["tourney_id"].rsplit("-", 1)[1]
    archive_name = (match.get("tourney_name") or "").strip()
    carried = event_names.get(code) or {}
    qualified = market_events.get(code) or {}
    location = str(carried.get("market_location", "") or qualified.get("market_location", ""))
    number = str(
        carried.get("market_event_number", "") or qualified.get("market_event_numbers", "")
    )
    for name, basis in (
        (
            str(carried.get("market_tournament", "")),
            "carried_tennis_data_name_matches_archive_name",
        ),
        (
            str(qualified.get("market_tournament", "")),
            "qualified_crosswalk_name_matches_archive_name",
        ),
    ):
        if name and archive_name and join.normalized(name) == join.normalized(archive_name):
            return {"ATP": number, "Location": location, "Tournament": name}, basis, ""
    if archive_name:
        warning = "" if carried or qualified else "no_carried_tennis_data_location_or_number"
        return (
            {"ATP": number, "Location": location, "Tournament": archive_name},
            "archive_event_name_exact_by_construction",
            warning,
        )
    name = str(carried.get("market_tournament", "") or qualified.get("market_tournament", ""))
    return (
        {"ATP": number, "Location": location, "Tournament": name},
        "carried_tennis_data_name_without_archive_name",
        "archive_event_carries_no_tourney_name",
    )


def wta_market_event_identity(
    match: Mapping[str, str],
    event_names: Mapping[str, Mapping[str, Any]],
) -> tuple[dict[str, str], str, str]:
    """The `Location`/`Tournament`/`WTA` of one synthesized WTA market row.

    The WTA join pairs a workbook block `(Location, Tournament)` with a crosswalk
    edition, and the carry-forward resolves it by that same pair (tier A) or by
    `Location` (tier B).  A synthesized row therefore carries the tennis-data identity
    the frozen crosswalk recorded for this mirror event's latest accepted edition.  An
    event the crosswalk never held gets the mirror's own name in both columns and a
    warning: it will resolve only if a tier-B location carry exists, and otherwise
    stays `edition_unresolved`, reported.
    """
    code = match["tourney_id"].rsplit("-", 1)[1]
    archive_name = (match.get("tourney_name") or "").strip()
    carried = event_names.get(code) or {}
    if carried:
        return (
            {
                "WTA": str(carried.get("market_event_number", "")),
                "Location": str(carried.get("market_location", "")),
                "Tournament": str(carried.get("market_tournament", "")),
            },
            "carried_tennis_data_identity_from_frozen_wta_crosswalk",
            "",
        )
    return (
        {"WTA": "", "Location": archive_name, "Tournament": archive_name},
        "archive_event_name_no_frozen_crosswalk_identity",
        "no_frozen_wta_crosswalk_edition_for_this_event_code",
    )


def wta_market_round(round_label: str, draw_size: Any) -> str | None:
    """The tennis-data round string for a mirror round label.

    The same inversion `join.round_agrees` performs (`R<n>` -> the ordinal of the
    bracket round it sits in, QF/SF/F fixed), restated because the WTA chain does not
    use MULTI01's ATP join.  `wta_sources.td_round_depth` reads these strings back
    relative to the block's deepest ordinal.
    """
    label = (round_label or "").strip()
    fixed = {"QF": "Quarterfinals", "SF": "Semifinals", "F": "The Final", "RR": "Round Robin"}
    if label in fixed:
        return fixed[label]
    top = bracket_top(draw_size)
    if top is None or not re.fullmatch(r"R\d+", label):
        return None
    index = (top // int(label[1:])).bit_length()
    return {1: "1st Round", 2: "2nd Round", 3: "3rd Round", 4: "4th Round"}.get(index)


def tennis_data_cover_index(
    market_rows: Sequence[Mapping[str, Any]], join: Any, *, by_location: bool = False
) -> dict[str, list[tuple[dt.date | None, str]]]:
    """Which (event, round) the season's retained tennis-data workbook already carries.

    The retained workbook covers every event through its own last date, and
    `prepare_panel` raises on the *second* market row that resolves to one archive row,
    so a round the workbook already carries must get no synthesized rows at all.  The
    unit is the round, not the match, deliberately: a per-match test has to compare
    player names across two sources and a single spelling difference then reads as a
    missing row and synthesizes a colliding one.
    """
    index: dict[str, list[tuple[dt.date | None, str]]] = defaultdict(list)
    for row in market_rows:
        key = join.normalized(str(row.get("Tournament", "")))
        index[key].append((row.get("market_date"), str(row.get("Round", "")).strip()))
        if by_location:
            # WTA: a sponsor rename between the carried season and the workbook must not
            # synthesize a second copy of a round the workbook already carries under the
            # new name; the location is the crosswalk's own primary event evidence.
            location = join.normalized(str(row.get("Location", "")))
            if location:
                index[f"location:{location}"].append(
                    (row.get("market_date"), str(row.get("Round", "")).strip())
                )
    return index


def _refuse(
    refused: list[dict[str, Any]],
    report: dict[str, Any],
    season: int,
    tour: str,
    match: Mapping[str, str],
    label: str,
    reason: str,
    detail: str,
) -> None:
    refused.append(_unmapped(season, tour, match["tourney_name"], 1, reason, detail))
    report["refused_per_event"][label] += 1


def synthesize_market_rows(
    added: Sequence[Mapping[str, str]],
    season: int,
    columns: Sequence[str],
    crosswalk_rows: Mapping[str, Mapping[str, str]],
    market_events: Mapping[str, Mapping[str, str]],
    join: Any,
    start_row: int,
    event_names: Mapping[str, Mapping[str, Any]] | None = None,
    covered: Mapping[str, list[tuple[dt.date | None, str]]] | None = None,
    tour: str = "ATP",
    event_end_dates: Mapping[str, dt.date] | None = None,
) -> tuple[list[dict[str, Any]], list[dict[str, Any]], list[dict[str, Any]], dict[str, Any]]:
    """One tennis-data-shaped row per added Sackmann match, prices blank.

    `prepare_panel` keeps only archive rows a *market* row resolves to, so a
    Wikipedia-only match needs a market row or it never reaches the panel.  The event's
    tennis-data `Tournament` comes from `market_event_identity` (`wta_market_event_identity`
    for WTA), which refuses nothing, and a match the retained workbook already carries
    gets no row at all.  Every price column is blank: `ps_missing = 1`, as the design
    declares for the window with no tennis-data prices.

    `Date` is the event's end (`event_end_dates`, keyed by generated `tourney_id`), the
    bound `generate_from_wikipedia` measured for the draw; a match with no end date is
    refused (`no_event_end_date`).  The round order is recorded as provenance only.
    """
    rows: list[dict[str, Any]] = []
    provenance: list[dict[str, Any]] = []
    refused: list[dict[str, Any]] = []
    event_names = event_names or {}
    covered = covered if covered is not None else {}
    event_end_dates = event_end_dates if event_end_dates is not None else {}
    report: dict[str, Any] = {
        "synthesized_per_event": defaultdict(int),
        "skipped_tennis_data_covered_per_event": defaultdict(int),
        "refused_per_event": defaultdict(int),
        "event_identity_basis": {},
        "event_identity_warnings": {},
    }
    row_number = start_row
    wta = tour.upper() == "WTA"
    for match in added:
        code = match["tourney_id"].rsplit("-", 1)[1]
        if wta:
            event, basis, warning = wta_market_event_identity(match, event_names)
        else:
            event, basis, warning = market_event_identity(match, market_events, event_names, join)
        label = f"{match['tourney_id']}|{match.get('tourney_name', '')}"
        report["event_identity_basis"][label] = basis
        if warning:
            report["event_identity_warnings"][label] = warning
        if wta:
            market_round = wta_market_round(match["round"], match["draw_size"])
        else:
            market_round = next(
                (
                    candidate
                    for candidate in MARKET_ROUND_CANDIDATES
                    if join.round_agrees(
                        {"round": match["round"], "draw_size": match["draw_size"]}, candidate
                    )
                ),
                None,
            )
        if market_round is None:
            _refuse(
                refused, report, season, tour, match, label,
                "no_market_round_string", f"round {match['round']} draw_size {match['draw_size']}",
            )  # fmt: skip
            continue
        start = dt.datetime.strptime(match["tourney_date"], "%Y%m%d").date()
        # The retained tennis-data workbook already carries every event through its own
        # last date, so a round it covers gets no synthesized rows.  The event is
        # recognised by either its archive name or the tennis-data name the carried
        # seasons recorded for it, inside the join's own -2..+21 day window around this
        # event's start.
        names = {join.normalized(event["Tournament"])}
        carried_name = str((event_names.get(code) or {}).get("market_tournament", ""))
        if carried_name:
            names.add(join.normalized(carried_name))
        if wta and event.get("Location"):
            names.add(f"location:{join.normalized(event['Location'])}")
        if any(
            existing_round == market_round
            and existing_date is not None
            and -2 <= (existing_date - start).days <= 21
            for name in names
            for existing_date, existing_round in covered.get(name, ())
        ):
            report["skipped_tennis_data_covered_per_event"][label] += 1
            continue
        winner = crosswalk_rows.get(str(match["winner_id"]))
        loser = crosswalk_rows.get(str(match["loser_id"]))
        if winner is None or loser is None:
            _refuse(
                refused, report, season, tour, match, label,
                "no_market_name_style_for_player", f"{match['winner_id']}/{match['loser_id']}",
            )  # fmt: skip
            continue
        offset = round_order(match["round"], match["draw_size"])
        end_date = event_end_dates.get(match["tourney_id"])
        if end_date is None:
            _refuse(
                refused, report, season, tour, match, label,
                "no_event_end_date", f"tourney_id {match['tourney_id']}",
            )  # fmt: skip
            continue
        anchor = end_date
        record: dict[str, Any] = dict.fromkeys(columns, "")
        record.update(
            {
                "Location": event["Location"],
                "Tournament": event["Tournament"],
                "Date": anchor.isoformat(),
                "Court": "",
                "Surface": match["surface"],
                "Round": market_round,
                "Best of": match["best_of"],
                "Winner": winner["market_style"],
                "Loser": loser["market_style"],
                "Comment": "Completed" if " RET" not in match["score"] else "Retired",
            }
        )
        if wta:
            # The WTA workbook's event-number column is `WTA` and its series column
            # `Tier`; a synthesized row carries the carried number and no tier.
            record.update({"WTA": event.get("WTA", ""), "Tier": ""})
        else:
            record.update({"ATP": event["ATP"], "Series": ""})
        record.update(
            market_season=season,
            market_source_path="",
            market_source_row=row_number,
            market_date=anchor,
        )
        rows.append(record)
        provenance.append(
            {
                "market_season": season,
                "market_source_row": row_number,
                "route": SYNTHESIZED_ROUTE,
                "tourney_id": match["tourney_id"],
                "round": match["round"],
                "winner_id": match["winner_id"],
                "loser_id": match["loser_id"],
                "price_source": "none_ps_missing",
                "market_date_basis": SYNTHESIZED_DATE_BASIS,
                "market_date": anchor.isoformat(),
                "round_order": "" if offset is None else offset,
            }
        )
        report["synthesized_per_event"][label] += 1
        row_number += 1
    for field in (
        "synthesized_per_event",
        "skipped_tennis_data_covered_per_event",
        "refused_per_event",
    ):
        report[field] = dict(sorted(report[field].items()))
    return rows, provenance, refused, report


WORKBOOK_EPOCH = dt.datetime(1980, 1, 1)


def write_market_workbook(
    path: Path, columns: Sequence[str], rows: Sequence[Mapping[str, Any]]
) -> str:
    """One `Data` sheet, header first, every cell from `_cell`.

    ZIP entry times and the created property are fixed. The B2 reconstruction found
    that openpyxl resets the modified property during save, so strict workbook byte
    determinism is not established. That residual docProps timestamp and its downstream
    source-hash columns are an explicit provenance difference (docs/EQUIVALENCE.md).
    Cell contents, styles and workbook parts are unchanged by timestamp normalization.
    """
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook = openpyxl.Workbook(write_only=True)
    workbook.properties.created = WORKBOOK_EPOCH
    workbook.properties.modified = WORKBOOK_EPOCH
    sheet = workbook.create_sheet("Data")
    sheet.append(list(columns))
    for row in rows:
        sheet.append([_cell(row, name) for name in columns])
    buffer = io.BytesIO()
    workbook.save(buffer)
    with (
        zipfile.ZipFile(buffer) as source,
        zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as target,
    ):
        for info in source.infolist():
            entry = zipfile.ZipInfo(info.filename, date_time=WORKBOOK_EPOCH.timetuple()[:6])
            entry.compress_type = zipfile.ZIP_DEFLATED
            target.writestr(entry, source.read(info.filename))
    return sha256(path)


def _cell(row: Mapping[str, Any], name: str) -> Any:
    """One workbook cell, with `Date` written as a real date cell.

    `join.read_market_workbook` takes `market_date` from
    `parse_date(value, "xlsx_date" if cell.is_date else <data type>)`, so a date written
    as *text* parses to `None`.  `record["Date"]` is already the canonicalized ISO
    string, so the date cell is written from the parsed `market_date` the reader itself
    produced.
    """
    if name == "Date" and isinstance(row.get("market_date"), dt.date):
        return row["market_date"]
    value = row.get(name, "")
    if isinstance(value, dt.date):
        return value.isoformat()
    if value is None:
        return ""
    return value


# ------------------------------------------------------------------ tar composition


def compose_tar(
    original: Path,
    tar_root: str,
    tour: str,
    span: Sequence[int],
    replacements: Mapping[int, bytes],
    destination: Path,
    *,
    allow_reserved: bool,
) -> dict[str, Any]:
    """A small tar carrying only the members `archive_panel` reads.

    That stage opens exactly `<tar_root>/<tour>/<tour>_players.csv` and
    `<tar_root>/<tour>/<tour>_matches_<year>.csv` for each year in its span, so the
    composed archive holds only those.  Every member that is not replaced is copied
    byte-for-byte and its sha256 recorded, so the composition can be checked against
    the mirror afterwards.  Member headers are fixed (mtime 0, uid/gid 0, mode 0644,
    no owner names) as the archive wrote them; the gzip header's own mtime is fixed at
    0 too (the archive let it float to the wall clock, which was the only
    non-reproducible byte in its output) and its name field is the destination's
    basename without `.gz`, as `tarfile.open(..., "w:gz")` writes it.
    """
    destination.parent.mkdir(parents=True, exist_ok=True)
    members: list[dict[str, Any]] = []
    with (
        destination.open("wb") as raw,
        gzip.GzipFile(destination.name.removesuffix(".gz"), "wb", 9, raw, mtime=0) as gz,
        tarfile.open(fileobj=gz, mode="w") as target,
        tarfile.open(original, "r:gz") as source,
    ):
        players = f"{tar_root}/{tour.lower()}/{tour.lower()}_players.csv"
        wanted = [players] + [
            f"{tar_root}/{tour.lower()}/{tour.lower()}_matches_{year}.csv" for year in span
        ]
        for member in wanted:
            year_match = re.search(r"_matches_(\d{4})\.csv$", member)
            year = int(year_match.group(1)) if year_match else None
            if year is not None:
                _guard(year, allow_reserved)
            info = source.getmember(member)
            if year is not None and year in replacements:
                payload = replacements[year]
                route = "bridge_extended"
            else:
                extracted = source.extractfile(info)
                if extracted is None:
                    raise ChainError(f"cannot read mirror member {member}")
                payload = extracted.read()
                route = "mirror_bytes_unchanged"
            entry = tarfile.TarInfo(member)
            entry.size = len(payload)
            entry.mtime = 0
            entry.mode = 0o644
            entry.uid = entry.gid = 0
            entry.uname = entry.gname = ""
            target.addfile(entry, io.BytesIO(payload))
            members.append(
                {
                    "member": member,
                    "route": route,
                    "bytes": len(payload),
                    "sha256": sha256_bytes(payload),
                }
            )
    return {
        "path": relative_to_root(destination, label="composed_archive"),
        "sha256": sha256(destination),
        "bytes": destination.stat().st_size,
        "members": members,
    }


# ------------------------------------------------------------------ driver


def write_annual_csv(path: Path, rows: Sequence[Mapping[str, Any]]) -> str:
    """Write a Sackmann annual member with the mirror's own `\\n` line terminator.

    `atomic_csv` uses `csv`'s default `\\r\\n`, so a season the bridge adds nothing to
    would come out one byte per line larger than the mirror member it replaced.  With
    `\\n` a season with no added rows reproduces the mirror member exactly.
    """
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_suffix(path.suffix + f".tmp.{os.getpid()}")
    with temporary.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(
            handle,
            fieldnames=list(SACKMANN_FIELDS),
            extrasaction="raise",
            lineterminator="\n",
        )
        writer.writeheader()
        for row in rows:
            writer.writerow(
                {
                    field: ("" if row.get(field) is None else str(row.get(field)))
                    for field in SACKMANN_FIELDS
                }
            )
        handle.flush()
        os.fsync(handle.fileno())
    os.replace(temporary, path)
    return sha256(path)


def market_style(name: str) -> str:
    parts = (name or "").strip().split()
    if len(parts) < 2:
        return name.strip()
    return f"{' '.join(parts[1:])} {parts[0][0]}."


def _wta_header_columns(carry_years: Sequence[int]) -> list[str]:
    """A WTA season with no workbook takes the header of the latest WTAODDS01 workbook
    the carried seasons have."""
    manifest_path = resolve_under_root(WTA_MARKET_MANIFEST, label="wta market manifest")
    odds_records = {
        int(record["year"]): record
        for record in json.loads(manifest_path.read_text(encoding="utf-8"))["records"]
        if record.get("tour") == "WTA" and record.get("year")
    }
    header_year = max(year for year in odds_records if year <= max(carry_years))
    header_path = resolve_under_root(odds_records[header_year]["path"], label="wta header workbook")
    columns, _ = read_wta_workbook(header_path, header_year)
    return columns


def _atp_header_columns(carry_years: Sequence[int]) -> list[str]:
    headers_path = resolve_under_root(MARKET_HEADERS_BY_YEAR, label="market headers")
    return list(json.loads(headers_path.read_text(encoding="utf-8"))[str(max(carry_years))])


def _manifest_record(
    *,
    wta: bool,
    season: int,
    workbook_path: Path,
    workbook_sha: str,
    market_route: str,
    market_path: Path | None,
    tennis_data: Mapping[str, Any] | None,
) -> dict[str, Any]:
    if wta:
        # `wta_market_join` opens `record["path"]` (workspace-relative) from a
        # WTAODDS01-shaped manifest and keys it by `tour`/`year`; the bridge's own
        # seasons are the only years WTAODDS01 does not carry.
        return {
            "tour": "WTA",
            "year": season,
            "path": relative_to_root(workbook_path, label="market workbook"),
            "sha256": workbook_sha,
            "bytes": workbook_path.stat().st_size,
            "format": "xlsx",
            "acquisition_route": market_route,
            "role": "bridge market workbook for the WTA02 join",
            "publisher_url": (tennis_data or {}).get("publisher_url", ""),
            "source_receipt": (tennis_data or {}).get("source_receipt", ""),
            "tennis_data_source": (
                None
                if market_path is None
                else {
                    "path": relative_to_root(market_path, label="tennis_data"),
                    "sha256": sha256(market_path),
                }
            ),
        }
    # `join.load_inputs` opens every retained record at `MARKET_DIR / record["retained_path"]`,
    # and `market_source_path` -- which `prepare_panel` and the pinned alias evidence key
    # on -- is derived from that same join.  The chain leaves `join.market_dir` at
    # MULTI01's own directory and only the *manifest* is the bridge's, so a carried record
    # resolves exactly as MULTI01 resolves it.  The bridge's own seasons are the only
    # records whose path changes, and they are seasons MULTI01 never had.
    base_market_dir = resolve_under_root(MARKET_ACQUISITION_DIR, label="market acquisition dir")
    return {
        "year": season,
        "retained_path": os.path.relpath(workbook_path, base_market_dir),
        "retained_sha256": workbook_sha,
        "retained_bytes": workbook_path.stat().st_size,
        "acquisition_route": market_route,
        "capture_timestamp": "",
        "expected_extension": "xlsx",
        "publisher_url": (tennis_data or {}).get("publisher_url", ""),
        "source_receipt": (tennis_data or {}).get("source_receipt", ""),
    }


def _merged_manifest(
    *, wta: bool, tour: str, seasons: Sequence[int], manifest_records: Sequence[dict[str, Any]]
) -> dict[str, Any]:
    if wta:
        # The WTA join reads a WTAODDS01-shaped manifest (`records` with `tour`, `year`,
        # `path`, `sha256`).  Every WTAODDS01 record is kept -- the join binds the frozen
        # document by hash elsewhere, so the copy here is for the bridge seasons only --
        # and a bridge season replaces any record of the same year.
        manifest_path = resolve_under_root(WTA_MARKET_MANIFEST, label="wta market manifest")
        base_manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
        kept = [
            row
            for row in base_manifest["records"]
            if not (row.get("tour") == "WTA" and row.get("year") and int(row["year"]) in seasons)
        ]
        return {
            "id": "WTA02-bridge-market-manifest",
            "base_manifest": {
                "path": relative_to_root(manifest_path, label="wta market manifest"),
                "sha256": sha256(manifest_path),
                "acquisition_id": base_manifest.get("acquisition_id"),
            },
            "scope": (
                f"WTA annual tennis-data workbooks: WTAODDS01's records unchanged, plus "
                f"bridge seasons {list(seasons)}; a bridge season is the retained RES2026 "
                "workbook carried unchanged, or a merge of it with synthesized "
                "Wikipedia-draw rows"
            ),
            "records": [*kept, *manifest_records],
            "bridge_records": [record["year"] for record in manifest_records],
        }
    manifest_path = resolve_under_root(MARKET_ACQUISITION_MANIFEST, label="market manifest")
    base_manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    kept = [row for row in base_manifest["records"] if int(row.get("year", 0)) not in seasons]
    return {
        "scope": (
            f"{tour} annual Tennis-Data files "
            f"{min(int(row['year']) for row in kept)}-{max(seasons)}; "
            "seasons beyond the MULTI01 span are bridge merges of the retained "
            "tennis-data workbook and synthesized Wikipedia-draw rows"
        ),
        "records": [*kept, *manifest_records],
        "bridge_records": [record["year"] for record in manifest_records],
    }


def build(config_path: Path, *, allow_reserved: bool) -> dict[str, Any]:
    document = read_config(config_path)
    plan = year_plan(document)
    section = document.get("bridge")
    if not isinstance(section, dict):
        raise ChainError("configuration has no bridge object")
    seasons = [int(year) for year in section["seasons"]]
    reserved = [year for year in seasons if year in RESERVED_YEARS]
    if reserved and not (allow_reserved and section.get("reserved_release_acknowledged") is True):
        raise ChainError(
            f"seasons {reserved} are inside the reserved window; pass --allow-reserved-years "
            "and set bridge.reserved_release_acknowledged = true after logging the exposure "
            "event in registries/exposure_log.jsonl"
        )
    archive = resolve_under_root(section["archive"]["path"], label="archive")
    require_hash(archive, section["archive"].get("sha256"), label="archive")
    tar_root = section["archive"]["tar_root"]
    tour = section.get("tour", "ATP").upper()
    wta = tour == "WTA"
    carry_years = [int(year) for year in section["carry_forward_years"]]
    output_dir = resolve_output_under_root(section["output_dir"], label="output_dir")
    policy = str(section.get("unresolved_event_policy", "refuse"))

    index = event_index(archive, tar_root, tour, carry_years, allow_reserved=allow_reserved)
    # Refuse to generate anything until every RES2026 event name of this tour resolves
    # to a carried mirror edition (policy `refuse`); under `quarantine` an unresolved
    # in-scope event is listed here and quarantined whole by `generate_from_wikipedia`.
    resolution_report = event_resolution_table(index, tour, policy)
    crosswalks = {tour: crosswalk_v2.build(archive, tour)}
    resolver = Resolver(crosswalks)
    if wta:
        join: Any = _WtaJoinShim()
        events_path = None
        market_events: dict[str, dict[str, str]] = {}
        names_path = WTA_MARKET_EVENT_NAMES_FILE
    else:
        join = join_module()
        events_path = resolve_under_root(
            section.get("market_event_crosswalk", DEFAULT_MARKET_EVENT_CROSSWALK),
            label="market_event_crosswalk",
        )
        _, event_rows = read_csv_rows(events_path)
        market_events = {}
        for row in sorted(event_rows, key=lambda item: int(item["season"])):
            tourney_id = row["archive_tourney_id"].strip()
            if "-" in tourney_id:
                market_events[tourney_id.rsplit("-", 1)[1]] = row
        names_path = MARKET_EVENT_NAMES_FILE
    # The tennis-data event identity the join itself recorded for every carried event,
    # not just the ones the qualified crosswalk had to adjudicate.
    carried_event_names_document = json.loads(names_path.read_text(encoding="utf-8"))
    carried_event_names = carried_event_names_document["events"]

    summary: dict[str, Any] = {
        "id": "WTA02-res2026-bridge" if wta else "CONFIRM2026-res2026-bridge",
        "year_plan": plan.as_document(),
        "tour": tour,
        "seasons": seasons,
        "carry_forward_years": carry_years,
        "unresolved_event_policy": policy,
        "archive": {"path": relative_to_root(archive, label="archive"), "sha256": sha256(archive)},
        "market_event_crosswalk": (
            None
            if events_path is None
            else {
                "path": relative_to_root(events_path, label="market_event_crosswalk"),
                "sha256": sha256(events_path),
            }
        ),
        "market_event_names": {
            "path": f"tennislab/sources/{names_path.name}",
            "sha256": sha256(names_path),
            "events": len(carried_event_names),
            "source": carried_event_names_document["source"],
        },
        "code": {
            "bridge": code_receipt(__name__),
            "crosswalk_v2": code_receipt("tennislab.panel.crosswalk_v2"),
            "join_candidates": None if wta else code_receipt("tennislab.panel.join"),
            "wta_sources": (
                {**code_receipt(__name__), "readers": "inlined from WTA01_event_map/wta_sources.py"}
                if wta
                else None
            ),
        },
        "seasons_detail": {},
        "outputs": {},
    }
    summary["res2026_event_resolution"] = resolution_report
    duplicates: list[dict[str, Any]] = []
    unmapped: list[dict[str, Any]] = []
    event_resolutions: dict[str, dict[str, Any]] = {}
    market_provenance: list[dict[str, Any]] = []
    replacements: dict[int, bytes] = {}
    manifest_records: list[dict[str, Any]] = []
    infobox_dir = (
        resolve_under_root(section["infobox_html_dir"], label="infobox_html_dir")
        if section.get("infobox_html_dir")
        else None
    )

    for season in seasons:
        sources = section["sources"][str(season)]
        wiki_paths = [
            resolve_under_root(item, label="wikipedia") for item in sources.get("wikipedia", [])
        ]
        wiki_rows = read_wikipedia(wiki_paths) if wiki_paths else []
        _payload, _header, mirror_rows = read_mirror_annual(
            archive, tar_root, tour, season, allow_reserved=allow_reserved
        )
        generated, _rejects, season_unmapped, season_resolutions = generate_from_wikipedia(
            wiki_rows, season, index, resolver, infobox_dir=infobox_dir
        )
        event_resolutions[str(season)] = season_resolutions
        event_end_dates = {
            record["generated_tourney_id"]: dt.date.fromisoformat(record["tournament_end_date"])
            for record in season_resolutions.values()
        }
        merged, season_duplicates = merge_annual(mirror_rows, generated, season)
        added = merged[len(mirror_rows) :]
        duplicates.extend(season_duplicates)
        unmapped.extend(season_unmapped)

        annual_path = output_dir / "sackmann" / f"{tour.lower()}_matches_{season}.csv"
        annual_sha = write_annual_csv(annual_path, merged)
        replacements[season] = annual_path.read_bytes()

        tennis_data = sources.get("tennis_data")
        market_rows: list[dict[str, Any]] = []
        columns: list[str] = []
        market_path: Path | None = None
        if tennis_data:
            market_path = resolve_under_root(tennis_data["path"], label="tennis_data")
            require_hash(market_path, tennis_data.get("sha256"), label="tennis_data")
            columns, market_rows = read_tennis_data(market_path, season, tour)
            for row in market_rows:
                market_provenance.append(
                    {
                        "market_season": season,
                        "market_source_row": row["market_source_row"],
                        "route": "tennis_data_passthrough",
                        "tourney_id": "",
                        "round": str(row.get("Round", "")),
                        "winner_id": "",
                        "loser_id": "",
                        "price_source": "tennis_data_workbook",
                        "market_date_basis": TENNIS_DATA_DATE_BASIS,
                        "market_date": (
                            "" if row.get("market_date") is None else row["market_date"].isoformat()
                        ),
                        "round_order": "",
                    }
                )
        if not columns:
            columns = _wta_header_columns(carry_years) if wta else _atp_header_columns(carry_years)
        crosswalk_rows = {
            str(row["winner_id"]): {"market_style": market_style(row["winner_name"])}
            for row in merged
        }
        crosswalk_rows.update(
            {
                str(row["loser_id"]): {"market_style": market_style(row["loser_name"])}
                for row in merged
            }
        )
        synthesized, synth_provenance, refused, synth_report = synthesize_market_rows(
            added,
            season,
            columns,
            crosswalk_rows,
            market_events,
            join,
            start_row=len(market_rows) + 2,
            event_names=carried_event_names,
            covered=tennis_data_cover_index(market_rows, join, by_location=wta),
            tour=tour,
            event_end_dates=event_end_dates,
        )
        market_provenance.extend(synth_provenance)
        unmapped.extend(refused)
        if not synthesized and market_path is not None:
            # A season the bridge has nothing to add to is carried, not rewritten:
            # rewriting it changed `market_source_sha256` and `market_source_path` on
            # every panel row of the season for no gain, and the acquired workbook is the
            # better provenance record anyway.
            workbook_path = market_path
            workbook_sha = sha256(market_path)
            market_route = "retained_workbook_carried_unchanged_no_synthesized_rows"
        else:
            workbook_path = output_dir / "market" / f"{season}.xlsx"
            workbook_sha = write_market_workbook(
                workbook_path, columns, [*market_rows, *synthesized]
            )
            market_route = "res2026_bridge_merge_tennis_data_plus_synthesized_wikipedia_rows"
        manifest_records.append(
            _manifest_record(
                wta=wta,
                season=season,
                workbook_path=workbook_path,
                workbook_sha=workbook_sha,
                market_route=market_route,
                market_path=market_path,
                tennis_data=tennis_data,
            )
        )
        summary["seasons_detail"][str(season)] = {
            "wikipedia_source_files": [
                relative_to_root(path, label="wikipedia") for path in wiki_paths
            ],
            "wikipedia_rows_read": len(wiki_rows),
            "generated_sackmann_rows": len(generated),
            "mirror_rows": len(mirror_rows),
            "rows_added_after_dedupe": len(added),
            "duplicates_dropped": len(season_duplicates),
            "duplicates_with_identical_keys": sum(
                1 for row in season_duplicates if row["keys_identical"]
            ),
            "annual_output": {
                "path": relative_to_root(annual_path, label="annual_output"),
                "sha256": annual_sha,
                "rows": len(merged),
            },
            "market_rows_passthrough": len(market_rows),
            "market_rows_synthesized": len(synthesized),
            "market_rows_synthesized_per_event": synth_report["synthesized_per_event"],
            "market_rows_skipped_tennis_data_covered": sum(
                synth_report["skipped_tennis_data_covered_per_event"].values()
            ),
            "market_rows_skipped_per_event": synth_report["skipped_tennis_data_covered_per_event"],
            "market_rows_refused_per_event": synth_report["refused_per_event"],
            "market_event_identity_basis": synth_report["event_identity_basis"],
            "market_event_identity_warnings": synth_report["event_identity_warnings"],
            "market_workbook": {
                "path": relative_to_root(workbook_path, label="market workbook"),
                "sha256": workbook_sha,
                "route": market_route,
            },
            # A changed or incomplete draw is a warning on the draw, not a quarantined
            # event, so the differences are reported per season.
            "draw_structure_warnings": {
                name: record["warnings"]
                for name, record in sorted(season_resolutions.items())
                if record.get("warnings")
            },
            "draws_with_rounds_pending": {
                name: record["rounds_pending"]
                for name, record in sorted(season_resolutions.items())
                if record.get("rounds_pending")
            },
        }

    span = list(range(int(section.get("panel_start_year", 2005)), plan.panel_end_year + 1))
    composed = compose_tar(
        archive,
        tar_root,
        tour,
        span,
        replacements,
        output_dir / "archive" / f"{tour.lower()}_panel_source.tar.gz",
        allow_reserved=allow_reserved,
    )
    summary["composed_archive"] = composed

    merged_manifest = _merged_manifest(
        wta=wta, tour=tour, seasons=seasons, manifest_records=manifest_records
    )
    summary["outputs"] = {
        "market/acquisition_manifest.json": atomic_json(
            output_dir / "market" / "acquisition_manifest.json", merged_manifest
        ),
        "market/market_row_provenance.csv": atomic_csv(
            output_dir / "market" / "market_row_provenance.csv",
            MARKET_PROVENANCE_FIELDS,
            market_provenance,
        ),
        "quarantine/unmatched_players.csv": atomic_csv(
            output_dir / "quarantine" / "unmatched_players.csv", UNMATCHED_FIELDS, resolver.rejected
        ),
        "quarantine/dropped_duplicates.csv": atomic_csv(
            output_dir / "quarantine" / "dropped_duplicates.csv", DUPLICATE_FIELDS, duplicates
        ),
        "quarantine/unmapped_events.csv": atomic_csv(
            output_dir / "quarantine" / "unmapped_events.csv", UNMAPPED_EVENT_FIELDS, unmapped
        ),
    }
    summary["event_resolutions_per_season"] = event_resolutions
    all_resolutions = [
        record
        for season_records in event_resolutions.values()
        for record in season_records.values()
    ]
    summary["draw_structure"] = {
        "draws_generated": len(all_resolutions),
        "draws_with_a_warning": sum(1 for record in all_resolutions if record.get("warnings")),
        "draws_with_draw_size_from_the_acquired_draw": sum(
            1 for record in all_resolutions if record.get("draw_size_from") == "acquired_draw"
        ),
        "draws_with_rounds_pending": sum(
            1 for record in all_resolutions if record.get("rounds_pending")
        ),
        "rounds_pending_total": sum(
            len(record.get("rounds_pending") or ()) for record in all_resolutions
        ),
        "warning_kind_counts": dict(
            sorted(
                Counter(
                    warning.split(":", 1)[0]
                    for record in all_resolutions
                    for warning in record.get("warnings") or ()
                ).items()
            )
        ),
    }
    summary["quarantine_counts"] = {
        "unmatched_player_name_rows": len(resolver.rejected),
        "distinct_unmatched_player_names": len({row["source_name"] for row in resolver.rejected}),
        "dropped_duplicates": len(duplicates),
        "dropped_duplicates_with_identical_keys": sum(
            1 for row in duplicates if row["keys_identical"]
        ),
        "unmapped_event_or_row_entries": len(unmapped),
        "unmapped_reasons": dict(sorted(Counter(row["reason"] for row in unmapped).items())),
    }
    summary["limits"] = [
        "Wikipedia is a labelled secondary source; a generated row's score and winner "
        "come from a bracket page written after the fact.",
        "A Wikipedia-only match has no per-match date: its archive row carries the "
        "tournament week start as the event anchor, and its synthesized market row -- "
        "which is where prepare_panel takes match_date from -- carries the event's END "
        "(the retained page's infobox completion date, else the last week start plus 7 "
        "days), WTA02 declared rule 4, so every match of a draw shares one date and no "
        "result is admitted before the event could have finished. Rest-day and workload "
        "features over those rows are computed on that inferred date.",
        "A draw whose bracket differs from the carried edition's keeps its own round "
        "labels and draw_size, derived from its own round column and order; the "
        "difference is a warning in event_resolutions_per_season, not a quarantine. "
        "Byes and unrecorded first-round matches are indistinguishable, so a draw "
        "missing part of its first round is accepted as a smaller first round.",
        "A synthesized market row has no price. Those rows carry ps_missing = 1 and no "
        "market reference; they are not a Pinnacle observation.",
        "Generated rows have blank serve counts, so count_block_status = missing_all and "
        "the SR02 dynamic state is not updated from them.",
    ]
    atomic_json(output_dir / "bridge_summary.json", summary)
    return summary


def generate_preview(config_path: Path, season: int, wikipedia: Sequence[Path]) -> dict[str, Any]:
    """Generate Sackmann rows from the given draw files and report them BEFORE dedupe.

    `self_test` and `build` both report only what survives the overlap rule, so a draw
    whose rows all duplicate mirror rows reports `rows_added_after_dedupe = 0` and says
    nothing about what was generated.  This mode reports the generated rows themselves:
    per draw, the structure derived from it, its round labels and counts, and its
    dedupe keys.  It refuses a reserved season, reads no market file and writes nothing.
    """
    _guard(season, False)
    document = read_config(config_path)
    section = document["bridge"]
    archive = resolve_under_root(section["archive"]["path"], label="archive")
    require_hash(archive, section["archive"].get("sha256"), label="archive")
    tour = section.get("tour", "ATP").upper()
    carry_years = [int(year) for year in section["carry_forward_years"]]
    index = event_index(
        archive, section["archive"]["tar_root"], tour, carry_years, allow_reserved=False
    )
    rows = read_wikipedia([resolve_under_root(item, label="wikipedia") for item in wikipedia])
    resolver = Resolver({tour: crosswalk_v2.build(archive, tour)})
    generated, rejected, unmapped, resolutions = generate_from_wikipedia(
        rows, season, index, resolver
    )
    by_draw: dict[str, list[dict[str, str]]] = defaultdict(list)
    for row in generated:
        by_draw[row["tourney_id"]].append(row)
    return {
        "season": season,
        "tour": tour,
        "wikipedia_files": [str(Path(item)) for item in wikipedia],
        "wikipedia_rows_read": len(rows),
        "generated_rows_before_dedupe": len(generated),
        "unmatched_player_rows": len(rejected),
        "unmapped_event_entries": unmapped,
        "event_resolutions": resolutions,
        "generated_per_draw": {
            tourney_id: {
                "rows": len(draw),
                "draw_size": draw[0]["draw_size"],
                "round_counts": dict(sorted(Counter(row["round"] for row in draw).items())),
                "match_num_span": [
                    min(int(row["match_num"]) for row in draw),
                    max(int(row["match_num"]) for row in draw),
                ],
                "dedupe_keys": sorted(
                    "|".join(
                        dedupe_key(
                            row["tourney_id"], row["round"], row["winner_id"], row["loser_id"]
                        )
                    )
                    for row in draw
                ),
            }
            for tourney_id, draw in sorted(by_draw.items())
        },
    }


# ------------------------------------------------------------------ 2024 self-test


def self_test(
    config_path: Path, year: int, sample: int = 20, seed: int = 20260911
) -> dict[str, Any]:
    """Round-trip `sample` mirror rows of `year` through the Wikipedia path.

    Builds a synthetic Wikipedia-format CSV from real mirror rows -- article-style
    player names (one deliberately carrying a `(tennis)` disambiguator), the page's
    round label rather than Sackmann's, a week label instead of a date, and Wikipedia's
    score rendering -- then asserts the regenerated Sackmann rows reproduce the mirror
    rows' keys.  The whole draw is fed in, not only the sampled rows, because the round
    alignment is a per-draw property.
    """
    import random

    if year in RESERVED_YEARS:
        raise ChainError(f"year {year} is inside the reserved outcome window")
    document = read_config(config_path)
    section = document["bridge"]
    archive = resolve_under_root(section["archive"]["path"], label="archive")
    tar_root = section["archive"]["tar_root"]
    tour = section.get("tour", "ATP").upper()
    _payload, _header, mirror_rows = read_mirror_annual(
        archive, tar_root, tour, year, allow_reserved=False
    )
    index = event_index(archive, tar_root, tour, [year - 1, year], allow_reserved=False)

    rng = random.Random(seed)
    eligible = [
        row
        for row in mirror_rows
        if row["surface"]
        and row["draw_size"].isdigit()
        and row["best_of"] in {"3", "5"}
        and row["winner_id"]
        and row["loser_id"]
        and find_edition(index, row["tourney_name"])[0] is not None
    ]
    chosen = rng.sample(eligible, sample)
    draws = {row["tourney_id"] for row in chosen}
    draw_rows = [row for row in mirror_rows if row["tourney_id"] in draws]

    # Mirror round label -> the page's ordered label, per draw, inverting align_rounds.
    page_label: dict[tuple[str, str], str] = {}
    week_label: dict[str, str] = {}
    for tourney_id in sorted(draws):
        rows = [row for row in mirror_rows if row["tourney_id"] == tourney_id]
        code = tourney_id.rsplit("-", 1)[1]
        edition = index[code]
        numbered = [label for label in edition["ordered_rounds"] if re.fullmatch(r"R\d+", label)]
        for position, label in enumerate(numbered, start=1):
            page_label[(tourney_id, label)] = f"R{position}"
        for label in ("QF", "SF", "F"):
            page_label[(tourney_id, label)] = label
        anchor = dt.datetime.strptime(rows[0]["tourney_date"], "%Y%m%d").date()
        week_label[tourney_id] = (
            (anchor.strftime("%b %-d") if tour == "WTA" else anchor.strftime("%-d %b"))
            if sys.platform != "win32"
            else anchor.isoformat()
        )

    synthetic: list[dict[str, str]] = []
    for position, row in enumerate(draw_rows):
        status = "completed"
        score = row["score"]
        if re.search(r"\bRET\b", score.upper()):
            status = "retired"
            score = re.sub(r"\s*RET\s*$", " 0-0 retired", score, flags=re.IGNORECASE)
        elif "W/O" in score.upper():
            status = "walkover"
            score = "w/o"
        label = page_label.get((row["tourney_id"], row["round"]))
        if label is None:
            continue
        winner = row["winner_name"]
        loser = row["loser_name"]
        # One row per draw carries a Wikipedia disambiguator, which the frozen Elo
        # crosswalk cannot match and crosswalk_v2 must.
        if position % 37 == 0:
            winner = f"{winner} (tennis)"
        synthetic.append(
            {
                "source": "wikipedia",
                "tour": tour,
                "tournament": row["tourney_name"],
                "round": label,
                "winner": winner,
                "loser": loser,
                "winner_display": row["winner_name"],
                "loser_display": row["loser_name"],
                "score": score.replace("-", "–"),
                "sets_played": "",
                "status": status,
                "week_label": week_label[row["tourney_id"]],
                "source_url": f"synthetic://{row['tourney_id']}",
            }
        )

    resolver = Resolver({tour: crosswalk_v2.build(archive, tour)})
    generated, rejected, unmapped, resolutions = generate_from_wikipedia(
        synthetic, year, index, resolver
    )
    _merged, dropped = merge_annual(mirror_rows, generated, year)

    compared = (
        "tourney_id", "round", "winner_id", "loser_id", "surface", "draw_size",
        "tourney_level", "tourney_date", "best_of", "tourney_name", "score",
    )  # fmt: skip
    by_key = {row["dedupe_key"]: row for row in dropped}
    checks: list[dict[str, Any]] = []
    for row in chosen:
        key = "|".join(
            dedupe_key(row["tourney_id"], row["round"], row["winner_id"], row["loser_id"])
        )
        entry = by_key.get(key)
        if entry is None:
            checks.append({"dedupe_key": key, "status": "NOT_GENERATED"})
            continue
        mismatches = {
            field: {"generated": entry[f"generated_{field}"], "mirror": entry[f"mirror_{field}"]}
            for field in compared
            if str(entry[f"generated_{field}"]) != str(entry[f"mirror_{field}"])
        }
        checks.append(
            {
                "dedupe_key": key,
                "status": "IDENTICAL" if not mismatches else "DIFFERS",
                "mismatches": mismatches,
            }
        )
    failures = [item for item in checks if item["status"] != "IDENTICAL"]
    # Stronger than the requested 20-row sample: every regenerated row in the fed draws
    # is compared, so a defect the sample misses still fails the test.
    all_mismatches: list[dict[str, Any]] = []
    field_counts: Counter[str] = Counter()
    classes: Counter[str] = Counter()

    def _strip(value: str) -> str:
        return " ".join(token for token in str(value).split() if token != "0-0")

    for entry in dropped:
        differing = {
            field: {"generated": entry[f"generated_{field}"], "mirror": entry[f"mirror_{field}"]}
            for field in compared
            if str(entry[f"generated_{field}"]) != str(entry[f"mirror_{field}"])
        }
        if not differing:
            continue
        field_counts.update(differing.keys())
        # The mirror renders an abandoned set inconsistently: a retirement is usually
        # "6-4 RET" but sometimes "6-4 0-0 RET".  A Wikipedia-rendered score cannot
        # distinguish the two, so `normalize_score` takes the majority convention and
        # drops the 0-0.  Such a difference is classified, not swept away: the score is
        # in FORBIDDEN_MODEL_COLUMNS and never a feature, and it moves only the market
        # score-agreement flag.
        kind = "other"
        if set(differing) == {"score"} and _strip(differing["score"]["generated"]) == _strip(
            differing["score"]["mirror"]
        ):
            kind = "abandoned_set_0_0_rendering_only"
        classes[kind] += 1
        all_mismatches.append(
            {"dedupe_key": entry["dedupe_key"], "class": kind, "mismatches": differing}
        )
    unexplained = [
        item for item in all_mismatches if item["class"] != "abandoned_set_0_0_rendering_only"
    ]
    return {
        "year": year,
        "tour": tour,
        "sampled_mirror_matches": len(chosen),
        "draws_fed": len(draws),
        "draw_rows_fed": len(draw_rows),
        "synthetic_wikipedia_rows": len(synthetic),
        "rows_with_disambiguator": sum(1 for row in synthetic if "(tennis)" in row["winner"]),
        "generated_sackmann_rows": len(generated),
        "duplicates_detected": len(dropped),
        "duplicates_with_identical_keys": sum(1 for row in dropped if row["keys_identical"]),
        "rows_added_after_dedupe": len(generated) - len(dropped),
        "unmatched_player_rows": len(rejected),
        "unmapped_event_entries": len(unmapped),
        "event_resolutions": resolutions,
        "compared_fields": list(compared),
        "sampled_checks_identical": sum(1 for item in checks if item["status"] == "IDENTICAL"),
        "failures": failures[:10],
        "all_generated_rows_compared": len(dropped),
        "all_generated_rows_with_identical_keys": len(dropped) - len(all_mismatches),
        "all_generated_rows_differing": len(all_mismatches),
        "differing_field_counts": dict(sorted(field_counts.items())),
        "difference_classes": dict(sorted(classes.items())),
        "differences_not_explained_by_score_rendering": len(unexplained),
        "all_row_mismatch_examples": all_mismatches[:10],
        "match_num_note": (
            "match_num is allocated by the bridge above the carried edition's maximum and "
            "is deliberately not compared; a duplicate row is dropped before it is used."
        ),
        "status": (
            "PASS" if not failures and len(chosen) == sample and not unexplained else "FAIL"
        ),
    }


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--config", type=Path, required=True)
    parser.add_argument("--allow-reserved-years", action="store_true")
    parser.add_argument("--self-test-year", type=int)
    parser.add_argument(
        "--event-self-check",
        action="store_true",
        help="resolve every references/RES2026.md event name against the carried mirror "
        "editions and exit; opens no draw file and no reserved season",
    )
    parser.add_argument("--sample", type=int, default=20)
    parser.add_argument(
        "--generate-preview-season",
        type=int,
        help="generate rows from --wikipedia for this season and report them before "
        "the overlap dedupe; refuses a reserved season",
    )
    parser.add_argument("--wikipedia", type=Path, nargs="+", default=())
    parser.add_argument("--report", type=Path)
    parser.add_argument("--dry-run", action="store_true")
    args = parser.parse_args(argv)

    if args.dry_run:
        document = read_config(args.config)
        plan = year_plan(document)
        section = document["bridge"]
        archive = resolve_under_root(section["archive"]["path"], label="archive")
        require_hash(archive, section["archive"].get("sha256"), label="archive")
        for season, sources in section["sources"].items():
            for item in sources.get("wikipedia", []):
                resolve_under_root(item, label=f"wikipedia {season}")
            if sources.get("tennis_data"):
                path = resolve_under_root(
                    sources["tennis_data"]["path"], label=f"tennis_data {season}"
                )
                require_hash(
                    path, sources["tennis_data"].get("sha256"), label=f"tennis_data {season}"
                )
        print(
            json.dumps(
                {
                    "status": "dry_run_ok",
                    "year_plan": plan.as_document(),
                    "seasons": section["seasons"],
                },
                sort_keys=True,
            )
        )
        return 0

    if args.event_self_check:
        document = read_config(args.config)
        section = document["bridge"]
        archive = resolve_under_root(section["archive"]["path"], label="archive")
        require_hash(archive, section["archive"].get("sha256"), label="archive")
        tour = section.get("tour", "ATP").upper()
        carry_years = [int(year) for year in section["carry_forward_years"]]
        index = event_index(
            archive, section["archive"]["tar_root"], tour, carry_years, allow_reserved=False
        )
        report = event_resolution_table(
            index, tour, str(section.get("unresolved_event_policy", "refuse"))
        )
        if args.report:
            atomic_json(args.report, report)
        print(json.dumps(report, indent=2, sort_keys=True))
        return 0

    if args.generate_preview_season is not None:
        if not args.wikipedia:
            parser.error("--generate-preview-season needs --wikipedia")
        report = generate_preview(args.config, args.generate_preview_season, args.wikipedia)
        if args.report:
            atomic_json(args.report, report)
        print(json.dumps(report, indent=2, sort_keys=True))
        return 0 if not report["unmapped_event_entries"] else 1

    if args.self_test_year is not None:
        report = self_test(args.config, args.self_test_year, sample=args.sample)
        if args.report:
            atomic_json(args.report, report)
        print(json.dumps(report, indent=2, sort_keys=True))
        return 0 if report["status"] == "PASS" else 1

    summary = build(args.config, allow_reserved=args.allow_reserved_years)
    print(
        json.dumps(
            {
                "seasons": summary["seasons"],
                "quarantine_counts": summary["quarantine_counts"],
                "composed_archive": summary["composed_archive"]["path"],
            },
            indent=2,
            sort_keys=True,
        )
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
