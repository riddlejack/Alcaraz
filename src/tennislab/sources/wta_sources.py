"""Shared loaders and normalizers for the WTA01 event crosswalk.

Ported from the archive's ``references/WTA01_event_map/wta_sources.py``. It reads the 18
retained tennis-data WTA annual workbooks (WTAODDS01, 2007-2024) and the Sackmann WTA
mirror match files for the same seasons out of the ARCHIVE01 snapshot tarball. Nothing
is repaired; rows are reported as found.

The WTA ``join`` and ``event_carry_forward`` stages declare this module in their config
by path and hash. That entry stays a declared binding -- it is still resolved and hashed
into the stage manifest -- but the code that runs is this module, imported by name.

What changed from the archive revision: the absolute ``ROOT`` and the two absolute input
paths computed from it become workspace-relative constants resolved through
``tennislab.chain.common.resolve_under_root`` at call time, and
``load_mirror_season``'s mutable default-argument cache becomes a module-level dict.
Every normalizer, loader, surname class and round-depth rule is unchanged.

Scope guard: seasons 2025 and 2026 are refused by every loader here.
"""

from __future__ import annotations

import csv
import io
import re
import tarfile
import unicodedata
from collections import Counter, defaultdict
from datetime import date, datetime
from typing import Any

import openpyxl
import xlrd

from tennislab.chain.common import resolve_under_root

TD_BASE = "data/raw/WTAODDS01/raw/annual"
ARCHIVE_TAR = (
    "data/raw/ARCHIVE01/snapshot/"
    "tennis-sackmann-archive-83733587353df8a41f2fd4f516147d5aa83f5a8d.tar.gz"
)
TAR_PREFIX = "tennis-sackmann-archive-83733587353df8a41f2fd4f516147d5aa83f5a8d"
SEASONS = list(range(2007, 2025))
FORBIDDEN_SEASONS = {2025, 2026}

TD_EXT = {y: ("xls" if y <= 2012 else "xlsx") for y in SEASONS}


def _check_season(season):
    if int(season) in FORBIDDEN_SEASONS:
        raise ValueError(f"season {season} is out of the authorized read scope")
    return int(season)


# ---------------------------------------------------------------- normalizing

# Tokens carried by sponsor/competition titles that never identify the city.
# Folded out before comparing a tennis-data event title to a mirror name.
_GENERIC = {
    "open",
    "international",
    "internationaux",
    "tennis",
    "championships",
    "championship",
    "classic",
    "cup",
    "trophy",
    "trophee",
    "ladies",
    "lady",
    "women",
    "womens",
    "wta",
    "tour",
    "tournament",
    "tournoi",
    "grand",
    "prix",
    "masters",
    "series",
    "of",
    "the",
    "de",
    "du",
    "des",
    "la",
    "le",
    "les",
    "and",
    "at",
    "presented",
    "by",
    "for",
    "championnats",
    "torneo",
    "copa",
    "gp",
    "event",
    "invitational",
    "finals",
    "final",
    "cupen",
    "bowl",
    "tennistournament",
    "indoors",
    "indoor",
    "openn",
}


def strip_accents(value):
    return "".join(
        ch for ch in unicodedata.normalize("NFKD", value) if not unicodedata.combining(ch)
    )


def norm_name(value):
    """Loose fold of an event name: accents, punctuation, case, whitespace."""
    if value is None:
        return ""
    text = strip_accents(str(value)).lower()
    text = text.replace("&", " and ")
    text = re.sub(r"[^a-z0-9]+", " ", text)
    return " ".join(text.split())


def name_tokens(value):
    """Normalized tokens with generic/sponsor-shaped words folded out."""
    return [t for t in norm_name(value).split() if t not in _GENERIC and not t.isdigit()]


def norm_core(value):
    return " ".join(name_tokens(value))


def norm_surface(value):
    if value is None:
        return ""
    return norm_name(value)


# ------------------------------------------------------------- player surname

_INITIALS = re.compile(r"^[a-z]{1,3}(\.[a-z]{1,3})*\.?$")


def td_surname(value):
    """Surname tokens of a tennis-data 'Surname X.' player string.

    Trailing initials tokens are dropped; the last remaining token is used as the
    comparison surname so that compound surnames ('Martinez Sanchez M.J.') align with
    the mirror's 'Maria Jose Martinez Sanchez'.
    """
    text = strip_accents(str(value or "")).lower()
    text = text.replace("-", " ").replace("'", "")
    toks = [t for t in re.split(r"[^a-z0-9.]+", text) if t]
    while toks and ("." in toks[-1] or (len(toks[-1]) <= 2 and _INITIALS.match(toks[-1]))):
        toks.pop()
    return toks[-1] if toks else ""


def mirror_surname(value):
    """Comparison surname of a mirror 'First Last' player string."""
    text = strip_accents(str(value or "")).lower()
    text = text.replace("-", " ").replace("'", "")
    toks = [t for t in re.split(r"[^a-z0-9]+", text) if t]
    return toks[-1] if toks else ""


# ------------------------------------------------------------------ td rounds

MIRROR_ROUND_DEPTH = {
    "F": 0,
    "SF": 1,
    "QF": 2,
    "R16": 3,
    "R32": 4,
    "R64": 5,
    "R128": 6,
}
_TD_ORDINAL = re.compile(r"^(\d+)(st|nd|rd|th)\s+round$")


def td_round_depth(round_value, max_ordinal):
    """Rounds-from-the-final depth of a tennis-data round label.

    'Nth Round' depth is read off the block's own deepest ordinal so that the same label
    means R32 in a 32 draw and R128 in a 128 draw.
    """
    text = norm_name(round_value)
    if text in ("the final", "final", "f"):
        return "F"
    if text in ("semifinals", "semifinal", "sf"):
        return "SF"
    if text in ("quarterfinals", "quarterfinal", "qf"):
        return "QF"
    if text in ("round robin", "roundrobin", "rr"):
        return "RR"
    m = _TD_ORDINAL.match(text)
    if m and max_ordinal:
        depth = (max_ordinal - int(m.group(1))) + 3
        for label, value in MIRROR_ROUND_DEPTH.items():
            if value == depth:
                return label
        return f"D{depth}"
    return text or ""


def td_max_ordinal(round_values):
    best = 0
    for value in round_values:
        m = _TD_ORDINAL.match(norm_name(value))
        if m:
            best = max(best, int(m.group(1)))
    return best


# ------------------------------------------------------------------ td reader


def _read_xls(path):
    book = xlrd.open_workbook(path)
    sheet = book.sheet_by_index(0)
    header = [str(sheet.cell_value(0, c)).strip() for c in range(sheet.ncols)]
    rows = []
    for r in range(1, sheet.nrows):
        values = []
        for c in range(sheet.ncols):
            cell = sheet.cell(r, c)
            value = cell.value
            if cell.ctype == xlrd.XL_CELL_DATE and value != "":
                try:
                    parts = xlrd.xldate_as_tuple(value, book.datemode)
                    value = date(parts[0], parts[1], parts[2])
                except Exception:  # noqa: BLE001 - the archive keeps the raw cell value
                    pass
            values.append(value)
        rows.append(values)
    return header, rows


def _read_xlsx(path):
    book = openpyxl.load_workbook(path, read_only=True, data_only=True)
    sheet = book[book.sheetnames[0]]
    it = sheet.iter_rows(values_only=True)
    header = [("" if v is None else str(v)).strip() for v in next(it)]
    rows = [list(r) for r in it]
    book.close()
    return header, rows


def _as_date(value):
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, str) and value.strip():
        for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d/%m/%y", "%m/%d/%Y"):
            try:
                return datetime.strptime(value.strip(), fmt).date()
            except ValueError:
                pass
    return None


def _blank(row):
    return all(v is None or (isinstance(v, str) and v.strip() == "") for v in row)


def _text(value):
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def load_td_season(season):
    """Tennis-data rows for one season as dicts of the identifying columns."""
    season = _check_season(season)
    path = resolve_under_root(f"{TD_BASE}/{season}.{TD_EXT[season]}", label="wta workbook")
    header, raw = (_read_xls if TD_EXT[season] == "xls" else _read_xlsx)(str(path))
    index = {h: i for i, h in enumerate(header) if h}
    out = []
    for number, row in enumerate(raw, start=2):
        if _blank(row):
            continue

        def col(name, row=row):
            i = index.get(name)
            return row[i] if i is not None and i < len(row) else None

        out.append(
            {
                "season": season,
                "row": number,
                "wta": _text(col("WTA")),
                "location": _text(col("Location")),
                "tournament": _text(col("Tournament")),
                "date": _as_date(col("Date")),
                "tier": _text(col("Tier")),
                "court": _text(col("Court")),
                "surface": _text(col("Surface")),
                "round": _text(col("Round")),
                "best_of": _text(col("Best of")),
                "winner": _text(col("Winner")),
                "loser": _text(col("Loser")),
                "comment": _text(col("Comment")),
            }
        )
    return out


# -------------------------------------------------------------- mirror reader

_MIRROR_CACHE: dict[int, list[dict[str, Any]]] = {}


def load_mirror_season(season):
    """Sackmann WTA mirror rows for one season, read from the ARCHIVE01 tar."""
    season = _check_season(season)
    if season in _MIRROR_CACHE:
        return _MIRROR_CACHE[season]
    member = f"{TAR_PREFIX}/wta/wta_matches_{season}.csv"
    archive = resolve_under_root(ARCHIVE_TAR, label="archive snapshot")
    with tarfile.open(archive, "r:gz") as tar:
        handle = tar.extractfile(member)
        if handle is None:
            raise FileNotFoundError(member)
        text = handle.read().decode("utf-8-sig")
    rows = list(csv.DictReader(io.StringIO(text)))
    for row in rows:
        row["season"] = season
    _MIRROR_CACHE[season] = rows
    return rows


def mirror_date(value):
    text = str(value or "").strip()
    if len(text) == 8 and text.isdigit():
        try:
            return date(int(text[:4]), int(text[4:6]), int(text[6:]))
        except ValueError:
            return None
    return None


# ------------------------------------------------- season surname canonicaliser


def td_surname_tokens(value):
    """All surname tokens of a tennis-data player string.

    Tennis-data always writes the initials as a token containing '.', so only those are
    dropped. Hyphenated surnames are split so that 'Smashnova-Pistolesi A.' can meet the
    mirror's 'Anna Smashnova'.
    """
    text = strip_accents(str(value or "")).lower().replace("'", "")
    parts = [p for p in re.split(r"\s+", text) if p]
    parts = [p for p in parts if "." not in p]
    tokens = []
    for part in parts:
        for tok in re.split(r"[^a-z0-9]+", part):
            if tok:
                tokens.append(tok)
    return tokens


def mirror_surname_tokens(value, td_token_universe):
    """Surname tokens of a mirror 'Given [Middle...] Family' player string.

    The last whitespace token is always taken as the family name. A middle token is only
    treated as part of a compound family name when tennis-data itself uses that token as
    a surname somewhere in the same season; this keeps multi-word given names ('Anna Lena
    Groenefeld') from being folded into the surname class.
    """
    text = strip_accents(str(value or "")).lower().replace("'", "")
    parts = [p for p in re.split(r"[^a-z0-9]+", text) if p]
    if not parts:
        return []
    if len(parts) == 1:
        return [parts[0]]
    tokens = [parts[-1]]
    for tok in parts[1:-1]:
        if tok in td_token_universe:
            tokens.append(tok)
    return tokens


class _Union:
    def __init__(self):
        self.parent = {}

    def find(self, item):
        self.parent.setdefault(item, item)
        root = item
        while self.parent[root] != root:
            root = self.parent[root]
        while self.parent[item] != root:
            self.parent[item], item = root, self.parent[item]
        return root

    def union(self, a, b):
        ra, rb = self.find(a), self.find(b)
        if ra == rb:
            return
        # deterministic: the lexicographically smaller token is the root
        if rb < ra:
            ra, rb = rb, ra
        self.parent[rb] = ra


def season_surname_canon(td_rows, mirror_rows, fuzzy_cutoff=0.85, fuzzy=True, opponents=True):
    """Season-wide surname classes shared by the two sources.

    Returns (td_player -> class, mirror_player -> class, fuzzy_union_log). Tokens of one
    player's surname are unioned together, so a source that carries only part of a
    compound surname still lands in the same class. A logged fuzzy pass then links
    tennis-data surnames that have no exact token in the mirror season (transliteration
    and married-name spellings).
    """
    import difflib

    td_players = {r["winner"] for r in td_rows} | {r["loser"] for r in td_rows}
    mirror_players = {r["winner_name"] for r in mirror_rows} | {
        r["loser_name"] for r in mirror_rows
    }

    td_tokens = {p: td_surname_tokens(p) for p in td_players}
    universe = {t for toks in td_tokens.values() for t in toks}
    mirror_tokens = {p: mirror_surname_tokens(p, universe) for p in mirror_players}

    uf = _Union()
    for toks in list(td_tokens.values()) + list(mirror_tokens.values()):
        toks = [t for t in toks if t]
        for tok in toks:
            uf.find(tok)
        for tok in toks[1:]:
            uf.union(toks[0], tok)

    mirror_roots = {uf.find(toks[0]) for toks in mirror_tokens.values() if toks}
    mirror_primary = sorted({toks[0] for toks in mirror_tokens.values() if toks})

    fuzzy_log = []
    for player, toks in sorted(td_tokens.items()) if fuzzy else []:
        if not toks:
            continue
        if uf.find(toks[0]) in mirror_roots:
            continue
        best = difflib.get_close_matches(toks[-1], mirror_primary, n=2, cutoff=fuzzy_cutoff)
        if len(best) == 1 or (
            len(best) == 2
            and difflib.SequenceMatcher(None, toks[-1], best[0]).ratio()
            > difflib.SequenceMatcher(None, toks[-1], best[1]).ratio()
        ):
            target = best[0]
            fuzzy_log.append(
                {
                    "td_player": player,
                    "td_surname": toks[-1],
                    "mirror_surname": target,
                    "ratio": round(difflib.SequenceMatcher(None, toks[-1], target).ratio(), 4),
                }
            )
            uf.union(toks[-1], target)

    td_canon = {p: (uf.find(t[0]) if t else "") for p, t in td_tokens.items()}
    mirror_canon = {p: (uf.find(t[0]) if t else "") for p, t in mirror_tokens.items()}

    opponent_log = (
        _link_by_opponents(td_rows, mirror_rows, td_canon, mirror_canon, uf) if opponents else []
    )
    td_canon = {p: (uf.find(t[0]) if t else "") for p, t in td_tokens.items()}
    mirror_canon = {p: (uf.find(t[0]) if t else "") for p, t in mirror_tokens.items()}
    return td_canon, mirror_canon, fuzzy_log + opponent_log


def _link_by_opponents(
    td_rows,
    mirror_rows,
    td_canon,
    mirror_canon,
    uf,
    min_residual=3,
    min_share=0.7,
    min_margin=0.2,
    rounds=3,
):
    """Link surname classes one source uses where the other uses another.

    A married, transliterated or mis-entered surname ('Groth' / 'Gajdosova') leaves a
    residual: matches this source attributes to a surname class that the other source
    does not. The residual opponent multisets of the two classes are compared; a link is
    proposed only when they overlap by at least ``min_share``, uniquely, by a clear
    margin, and the two names share a first initial. This is season-wide player evidence
    computed over surname classes both sources already agree on; it never uses event
    identity.

    The links are a key-folding aid for event matching. They are NOT a player identity
    claim: a source that writes a different player's name under the same surname produces
    the same residual signature.
    """
    log = []
    for _ in range(rounds):
        td_canon = {p: uf.find(c) if c else "" for p, c in td_canon.items()}
        mirror_canon = {p: uf.find(c) if c else "" for p, c in mirror_canon.items()}

        td_opp, mi_opp = defaultdict(Counter), defaultdict(Counter)
        td_members, mi_members = defaultdict(set), defaultdict(set)
        for row in td_rows:
            w, l = td_canon.get(row["winner"], ""), td_canon.get(row["loser"], "")  # noqa: E741
            td_members[w].add(row["winner"])
            td_members[l].add(row["loser"])
            td_opp[w][l] += 1
            td_opp[l][w] += 1
        for row in mirror_rows:
            w = mirror_canon.get(row["winner_name"], "")
            l = mirror_canon.get(row["loser_name"], "")  # noqa: E741
            mi_members[w].add(row["winner_name"])
            mi_members[l].add(row["loser_name"])
            mi_opp[w][l] += 1
            mi_opp[l][w] += 1

        def residual(a, b):
            out = Counter()
            for key, count in a.items():
                left = count - b.get(key, 0)
                if left > 0:
                    out[key] = left
            return out

        td_resid = {c: residual(td_opp[c], mi_opp.get(c, Counter())) for c in td_opp}
        mi_resid = {c: residual(mi_opp[c], td_opp.get(c, Counter())) for c in mi_opp}
        td_side = sorted(c for c, r in td_resid.items() if c and sum(r.values()) >= min_residual)
        mi_side = sorted(c for c, r in mi_resid.items() if c and sum(r.values()) >= min_residual)

        made = []
        used = set()
        for t in td_side:
            a = td_resid[t]
            if len(td_members[t]) > 2:
                continue
            initials = {td_first_initial(p) for p in td_members[t]}
            scored = []
            for m in mi_side:
                if m == t or m in used or len(mi_members[m]) > 2:
                    continue
                if not (initials & {mirror_first_initial(p) for p in mi_members[m]}):
                    continue
                b = mi_resid[m]
                inter = sum(min(c, b.get(k, 0)) for k, c in a.items())
                denom = min(sum(a.values()), sum(b.values()))
                scored.append((inter / denom if denom else 0.0, inter, m))
            scored.sort(reverse=True)
            if not scored or scored[0][0] < min_share:
                continue
            if len(scored) > 1 and scored[0][0] - scored[1][0] < min_margin:
                continue
            value, inter, m = scored[0]
            made.append(
                {
                    "td_player": "|".join(sorted(td_members[t])),
                    "mirror_player": "|".join(sorted(mi_members[m])),
                    "td_surname": t,
                    "mirror_surname": m,
                    "ratio": round(value, 4),
                    "evidence": (
                        f"residual_opponent_overlap {inter}/"
                        f"{min(sum(a.values()), sum(mi_resid[m].values()))}"
                    ),
                }
            )
            used.add(m)
            uf.union(t, m)
        log.extend(made)
        if not made:
            break
    return log


# ------------------------------------------------------- relative round depth

# Canonical distance from the final, deep rounds first. Round robin sits below every
# knockout round; BR is the Olympic bronze play-off.
MIRROR_ROUND_ORDER = {
    "F": 0,
    "BR": 1,
    "SF": 2,
    "QF": 3,
    "R16": 4,
    "R32": 5,
    "R64": 6,
    "R128": 7,
    "RR": 100,
    "ER": 101,
    "Q1": 102,
    "Q2": 102,
    "Q3": 102,
}


def td_round_order(round_value, max_ordinal):
    """Canonical distance-from-the-final order of a tennis-data round label."""
    text = norm_name(round_value)
    if text in ("the final", "final", "f"):
        return 0
    if text in ("semifinals", "semifinal", "sf"):
        return 2
    if text in ("quarterfinals", "quarterfinal", "qf"):
        return 3
    if text in ("round robin", "roundrobin", "rr"):
        return 100
    m = _TD_ORDINAL.match(text)
    if m and max_ordinal:
        return 3 + (max_ordinal - int(m.group(1)) + 1)
    return 200


def relative_depths(order_values):
    """Map each distinct round order value to its rank counted from the final.

    Aligning on the rounds each edition actually contains absorbs the labelling
    difference where the mirror starts a 32 draw at 'R64' while tennis-data calls the
    same matches '1st Round'.
    """
    return {value: rank for rank, value in enumerate(sorted(set(order_values)))}


# -------------------------------------------------------------- name initials


def td_first_initial(value):
    text = strip_accents(str(value or "")).lower()
    for tok in reversed([p for p in re.split(r"\s+", text) if p]):
        if "." in tok:
            letters = [c for c in tok if c.isalpha()]
            if letters:
                return letters[0]
    return ""


def mirror_first_initial(value):
    text = strip_accents(str(value or "")).lower()
    parts = [p for p in re.split(r"[^a-z0-9]+", text) if p]
    return parts[0][0] if parts else ""
